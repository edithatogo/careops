from pathlib import Path
import sys
import hashlib
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from docx import Document
from upload_gate import inspect_docx, validate_manifest

def test_text_and_disguised_text_are_blocked(tmp_path):
    p=tmp_path/'draft.txt'; p.write_text('Draft')
    assert inspect_docx(p)
    p=p.rename(tmp_path/'draft.docx')
    assert inspect_docx(p)

def test_styled_document_is_not_plain(tmp_path):
    p=tmp_path/'draft.docx';d=Document();d.add_heading('Draft');d.save(p)
    assert inspect_docx(p)

def test_plain_docx_still_needs_current_complete_review(tmp_path):
    p=tmp_path/'draft.docx';d=Document();d.add_paragraph('Draft');d.save(p)
    assert not inspect_docx(p)
    digest=hashlib.sha256(p.read_bytes()).hexdigest()
    e={'path':p.name,'sha256':digest,'format_route':'plain_docx'}
    m={'artifacts':[e]}
    assert validate_manifest(m,tmp_path)
    e['review']={'sha256':digest,'scores':dict(accuracy=300,purpose_status=200,format=200,readability=150,release=150),'hard_stops':[],**{k:'synthetic evidence' for k in ('visual_pages_evidence','source_fidelity_evidence','release_evidence','independent_reviewer')}}
    assert not validate_manifest(m,tmp_path)
    d.add_paragraph('Changed');d.save(p)
    assert any('changed' in x for x in validate_manifest(m,tmp_path))

def test_normal_style_and_paragraph_decoration_are_blocked(tmp_path):
    from docx.oxml import OxmlElement
    p=tmp_path/'draft.docx';d=Document();d.styles['Normal'].font.bold=True;d.add_paragraph('Draft');d.save(p)
    assert inspect_docx(p)
    d=Document();p1=d.add_paragraph('Draft');p1._p.get_or_add_pPr().append(OxmlElement('w:shd'));d.save(p)
    assert inspect_docx(p)

def test_fake_package_and_malformed_manifest(tmp_path):
    from zipfile import ZipFile
    p=tmp_path/'draft.docx'
    with ZipFile(p,'w') as z:z.writestr('word/document.xml','<document/>')
    assert inspect_docx(p)
    for m in ({'artifacts':[None]}, {'artifacts':'bad'}, {'artifacts':[{'path':p.name,'review':42}]}):
        assert validate_manifest(m,tmp_path)

def test_duplicate_manifest_and_missing_active_manifest(tmp_path):
    p=tmp_path/'draft.docx';d=Document();d.add_paragraph('Draft');d.save(p)
    e={'path':p.name}
    assert any('duplicate' in x for x in validate_manifest({'artifacts':[e,e]},tmp_path))

def test_document_defaults_cannot_hide_decoration(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    p=tmp_path/'draft.docx';d=Document();d.add_paragraph('Draft')
    defaults=d.styles.element.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    defaults.append(OxmlElement('w:b'));d.save(p)
    assert inspect_docx(p)
