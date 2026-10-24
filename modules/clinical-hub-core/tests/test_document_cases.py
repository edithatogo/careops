"""Execute portable fictional cases through unchanged preparation helpers."""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path

from docx import Document
import pytest
from scripts import prepare_meeting_draft as agenda
from scripts import prepare_minutes_draft as minutes
from scripts.preflight_pack import validate as preflight

ROOT = Path(__file__).resolve().parents[1]
CASES = ROOT / 'tests/fixtures/document_cases'


def load(name):
    return json.loads((CASES / name).read_text())


def minutes_input(name):
    data = load(name)
    # Test preparation only: portable source paths become exact local bindings.
    for source in data['sources']:
        source['path'] = str(ROOT / source['path'])
    return data


@pytest.mark.parametrize('name,count,exceptions', [('empty', 0, 0), ('single', 1, 0), ('repeating', 12, 2)])
def test_agenda_pack_variants_preserve_counts_bodies_links_and_holds(name, count, exceptions):
    data = load(f'agenda-{name}.json'); pack = load(f'pack-{name}.json')
    original = deepcopy((data, pack))
    doc = Document(BytesIO(agenda.build_bytes(data)))
    text = '\n'.join(p.text for p in doc.paragraphs)
    links = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith('/hyperlink')]
    assert len(data['items']) == count
    assert [p.text for p in doc.paragraphs if p.text.startswith('DCASE-A')] == [f"{i['id']}: {i['title']}" for i in data['items']]
    assert len(links) == count
    assert text.count('Content unresolved — confirm before issue.') == (3 if count == 0 else 2)
    for item, paper in zip(data['items'], pack['papers'], strict=True):
        assert item['body'] in text
        assert item['source'] == paper['source'] and item['source'] in links
    result = preflight(pack)
    assert result['paper_count'] == count and len(result['exceptions']) == exceptions
    assert result['ready_for_manual_review'] == (exceptions == 0)
    if name == 'repeating':
        assert {(e['paper'], e['check']) for e in result['exceptions']} == {('DCASE-P02', 'review'), ('DCASE-P02', 'length')}
        assert 'Fictional long paragraph 40:' in text
    assert (data, pack) == original


def test_explicit_mapping_does_not_join_on_similar_ids_or_invent_notes():
    manifest = load('cases.json'); data = load('agenda-repeating.json'); pack = load('pack-repeating.json')
    assert manifest['occurrence_id'] == pack['meeting_id'] == 'DCASE-OCC-01'
    assert load('minutes-register.json')['meeting_id'] == 'DCASE-OCC-01'
    assert load('minutes-notes.json')['meeting_id'] == manifest['minutes_baseline']['meeting_id'] == 'F08-DEMO-01'
    mappings = manifest['mapping']
    assert len({m['agenda_item_id'] for m in mappings}) == len({m['paper_id'] for m in mappings}) == 12
    for item, paper, mapping in zip(data['items'], pack['papers'], mappings, strict=True):
        assert mapping['agenda_item_id'] == item['id']
        assert mapping['paper_id'] == paper['id']
        assert mapping['paper_link'] == paper['source'] == item['source']
    assert [m['minutes_agenda_ref'] for m in mappings[:4]] == ['A-1', 'A-2', 'A-3', 'A-4']
    assert [m['notes_ref'] for m in mappings[:4]] == ['N-1', 'N-2', 'N-3', None]
    assert all(m['minutes_agenda_ref'] is None and m['notes_ref'] is None for m in mappings[4:])
    assert [(m['record_id'], m['agenda_item_id']) for m in manifest['action_mapping']] == [
        ('F-01', 'DCASE-A01'), ('F-02', 'DCASE-A02'), ('F-04', None), ('F-05', 'DCASE-A03')]
    assert {m['record_id'] for m in manifest['action_mapping']} == {r['id'] for r in load('minutes-register.json')['records']}


def test_register_preserves_held_pending_and_closed_distinctions(tmp_path):
    data = minutes_input('minutes-register.json'); original = deepcopy(data)
    selected, excluded, _ = minutes.select(data)
    expected = load('cases.json')['register_expected']
    assert [r['id'] for r in selected] == expected['selected_ids'] == ['F-01', 'F-02']
    assert [r['id'] for r, _ in excluded] == expected['excluded_ids'] == ['F-04', 'F-05']
    text = '\n'.join(p.text for p in Document(minutes.build(data, tmp_path/'register.docx')).paragraphs)
    assert 'Closure remains awaiting acceptance' in text
    assert 'DCASE-REGISTER#F-04:closure-accepted' in text
    assert 'Recorded snapshot status: Held candidate' in text
    for row in data['records']:
        for field in ('conditions', 'dissent', 'dependencies'):
            assert row[field] in text
    assert [r['id'] for r in data['records'] if r['state'] == 'candidate'] == expected['held_ids']
    assert [r['id'] for r in data['records'] if r['state'] == 'closed'] == expected['closed_ids']
    assert [r['id'] for r in selected if r['state'] == 'closure_pending'] == expected['closure_pending_ids']
    assert data == original


def test_pinned_minutes_leave_missing_and_off_agenda_content_visible(tmp_path):
    data = minutes_input('minutes-notes.json')
    doc = Document(minutes.build(data, tmp_path/'notes.docx'))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert sum(p.text.startswith('Agenda A-') for p in doc.paragraphs) == 4
    assert len(data['minutes_notes']['entries']) == 5
    for required in ('authority unverified', 'not commissioned, assigned or issued',
                     'No source-linked note supplied', 'Off-agenda concerns — unresolved',
                     'F08-NOTES-01#F08-DEMO-01:N-4'):
        assert required in text
    assert 'fixture://F08-DEMO-01/agenda/revision-1' in text


@pytest.mark.parametrize('kind', ['agenda', 'register', 'notes'])
def test_regeneration_refuses_to_overwrite_saved_human_edits(kind, tmp_path):
    data = load('agenda-single.json') if kind == 'agenda' else minutes_input(f'minutes-{kind}.json')
    build = agenda.build if kind == 'agenda' else minutes.build
    output = tmp_path / f'{kind}.docx'
    build(data, output)
    edited = Document(output); edited.add_paragraph(load('cases.json')['edit_marker']); edited.save(output)
    before = output.read_bytes()
    with pytest.raises(FileExistsError):
        build(data, output)
    assert output.read_bytes() == before
    fresh = deepcopy(data)
    if kind == 'agenda': fresh['items'][0]['body'] += '\nFictional source revision 2.'
    elif kind == 'register': fresh['sources'][0]['revision'] = 'Fictional register revision 2'
    else:
        # Pinned F08 source revision cannot be silently refreshed.
        fresh['minutes_notes']['notes_revision'] = 'Fictional notes revision 2'
        with pytest.raises(ValueError): build(fresh, tmp_path/'fresh.docx')
        assert not (tmp_path/'fresh.docx').exists()
        return
    build(fresh, tmp_path/'fresh.docx')
    assert output.read_bytes() == before
    assert load('cases.json')['edit_marker'] in '\n'.join(p.text for p in Document(output).paragraphs)
