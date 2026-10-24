"""Connected fictional meeting rehearsal through the actual reconciled helpers.

Explicit source mappings join the stages; no cloud operation, inferred approval,
new register or live mail path is exercised. Outputs are disposable test documents.
"""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import uuid

from docx import Document
import pytest

from scripts import prepare_meeting_draft as agenda
from scripts import prepare_minutes_draft as minutes
from scripts.f10_source_adapter import register_source_snapshot, carry_identity_key
from scripts.optional_action_reference import propose as action_reference
from scripts.preflight_pack import validate as preflight
from scripts.preparation_cycle import next_cycle
from scripts.reviewed_action_comparison import action_discrepancies
from scripts.summarise_feedback import summarise
from scripts.meeting_commit_preview import preview
from scripts.f15_resolution_lifecycle import propose_lifecycle
from scripts.ddms_reviewed_contribution import prepare_review, reconcile_review
from scripts.f09_isolated_writer import plan_action_register_recovery, reconcile_selective_recovery
from tests.test_f10_register_source import config as source_config
from tests.test_f09_register_recovery import config as recovery_config
from tests.test_meeting_commit_preview import case as outcome_case
from tests.test_f15_resolution_lifecycle import state as lifecycle_state, req, response, closure
from tests.test_ddms_reviewed_contribution import fixture as contribution_fixture

ROOT = Path(__file__).resolve().parents[1]
CASES = ROOT / 'tests/fixtures/document_cases'


def fixture(tmp_path):
    data = json.loads((CASES/'minutes-register.json').read_text())
    replacements = {row['id']: 'FICTIONAL-'+row['id'] for row in data['records']}
    original_source = (CASES/'register.txt').read_text()
    for old, new in replacements.items():
        original_source = original_source.replace(old, new)
    source = tmp_path/'source.txt'; source.write_text(original_source)
    data['sources'][0].update(path=str(source), sha256=sha256(source.read_bytes()).hexdigest())
    cfg = source_config(); prototype = cfg['rows'][0]
    rows = []
    for index, record in enumerate(data['records'], 1):
        old = record['id']; record['id'] = replacements[old]
        if record['closure_evidence_ref']:
            record['closure_evidence_ref'] = record['closure_evidence_ref'].replace(old, record['id'])
        row = deepcopy(prototype)
        row.update(Id=index, GUID=str(uuid.UUID(int=index)), LegacyactionID=record['id'],
                   Title='FICTIONAL '+record['text'], Sourcereference=data['sources'][0]['id']+'#'+record['id'],
                   Currentstatus='held' if record['state']=='candidate' else record['state'],
                   Duedate=record['due_date'], Reviewdate=record['review_date'],
                   Sourcehistory=record['conditions']+' | '+record['dissent'],
                   Modified=data['sources'][0]['snapshot'])
        rows.append(row)
    cfg['rows'] = rows
    cfg['selection'].update(meeting_id=data['meeting_id'], meeting_date=data['meeting_date'],
                            rows=[{k:row[k] for k in ('Id','GUID','OData__UIVersionString','LegacyactionID')} for row in rows])
    cfg['meeting'].update(meeting_id=data['meeting_id'], meeting_date=data['meeting_date'])
    return cfg, data


def test_source_to_pack_minutes_review_follow_up_and_next_preparation(tmp_path):
    cfg, data = fixture(tmp_path); before = deepcopy((cfg, data))
    snapshot = register_source_snapshot(**cfg)
    selected, excluded, sources = minutes.select(data)
    assert [r['id'] for r in selected] == ['FICTIONAL-F-01','FICTIONAL-F-02']
    assert {r['id'] for r,_ in excluded} == {'FICTIONAL-F-04','FICTIONAL-F-05'}
    # Descriptive source readiness never supplies commissioning or ownership.
    assert all(not row['issued_claim'] and not row['commissioned_claim'] for row in snapshot['records'])
    native = {row['record_id']: row for row in snapshot['records']}
    pack = {'meeting_id':snapshot['meeting_id'], 'pack_source':'fixture://cycle/pack-v1','papers':[]}
    agenda_data = {'meeting_date':snapshot['meeting_date'], 'items':[]}
    for row in selected:
        link = 'https://example.test/cycle/'+row['id']+'?version=1'
        pack['papers'].append(dict(id=row['id'],title=row['text'],source=link,version='1',reviewed=True,
            reviewed_source=link,reviewed_version='1',format='brief',pages=1,purpose='discussion'))
        agenda_data['items'].append(dict(id=row['id'],title=native[row['id']]['native_business_fields']['Title'],
            body=native[row['id']]['source_text'],source=link,reviewed=True,purpose='discussion',section='main'))
    assert preflight(pack)['ready_for_manual_review']
    rendered = Document(BytesIO(agenda.build_bytes(agenda_data)))
    assert all(item['body'] in '\n'.join(p.text for p in rendered.paragraphs) for item in agenda_data['items'])
    output = minutes.build(data, tmp_path/'minutes.docx')
    text = '\n'.join(p.text for p in Document(output).paragraphs)
    assert all(row['id'] in text and row['conditions'] in text and row['dissent'] in text for row in data['records'])
    assert 'Closure remains awaiting acceptance' in text
    assert action_discrepancies([{'id':r['id'],'text':r['text']} for r in selected],
                                [{'id':r['id'],'text':r['text']} for r in selected]) == []
    d,s,current,candidate,review = outcome_case()
    identity = selected[0]['id']
    # The review is pinned to the actual generated minutes bytes, not the register's revision label.
    revision = 'sha256:' + sha256(output.read_bytes()).hexdigest()
    assert selected[0]['source_id'] in sources and selected[0]['text'] in text
    current.update(item_id=identity, version=native[identity]['snapshot_version'],wording=selected[0]['text'])
    s.update(source_ref='FICTIONAL-MINUTES-CYCLE', revision=revision,passage_ref='FICTIONAL-MINUTES-CYCLE#'+identity,
             passage_text=selected[0]['text'],conditions=[selected[0]['conditions']],dissent=[selected[0]['dissent']])
    candidate.update(item_id=identity,source_ref=s['source_ref'],source_revision=revision,
                     before_version=current['version'],passage_ref=s['passage_ref'],wording=s['passage_text'],
                     conditions=s['conditions'],dissent=s['dissent'])
    candidate.update(kind='action',commission=dict(explicit=True,authority_ref='FICTIONAL-COMMISSION',
        meeting_purpose='commission',owner_accepted=True,owner_ref='FICTIONAL-ACCEPTED-OWNER',
        owner_authority_ref='FICTIONAL-OWNER-AUTHORITY',owner_acceptance_ref='FICTIONAL-OWNER-ACCEPTANCE',
        date_agreed=True,date_kind='due',date_value=data['meeting_date'],date_authority_ref='FICTIONAL-DATE-AGREEMENT'))
    proposal = preview(destination=d,source=s,current=current,candidate=candidate,review=review,as_of_date=data['meeting_date'])
    assert proposal['state']=='APPROVED_PROPOSAL' and proposal['committed'] is False
    assert proposal['after']['conditions']==[selected[0]['conditions']]
    assert proposal['source']['revision']==revision
    edited_minutes=Document(output);edited_minutes.add_paragraph('Independent fictional human correction');edited_minutes.save(output)
    updated_source=dict(s,revision='sha256:'+sha256(output.read_bytes()).hexdigest())
    stale_review=preview(destination=d,source=updated_source,current=current,candidate=candidate,review=review,as_of_date=data['meeting_date'])
    assert stale_review['state']=='BLOCKED'
    assert 'source_version_mismatch' in {entry['code'] for entry in stale_review['blockers']}
    # Reference an already observed agenda row; never create a second action register.
    fields={'Title':agenda_data['items'][0]['title'],'Question_x0020_or_x0020_help_x00':agenda_data['items'][0]['body'],
            'Meeting_x0020_date':snapshot['meeting_date']}
    key=carry_identity_key(snapshot['source_id'],identity,snapshot['meeting_id'])
    existing=[dict(fields,CarryForwardKey=key,Id=7,OData__UIVersionString='1.0',**{'odata.etag':'"1"'})]
    envelope=dict(contract_version=1,source_id=snapshot['source_id'],record_id=identity,meeting_id=snapshot['meeting_id'],
                  revision=snapshot['source_revision'],operation_key=proposal['operation_key'],fields=fields)
    reference=action_reference(existing,envelope,expected_revision=snapshot['source_revision'])
    assert reference['status']=='EXISTING_MATCH' and reference['write_proposed'] is False
    assert action_reference(existing,envelope,expected_revision=snapshot['source_revision'],prior=reference['receipt'])['status']=='duplicate'
    edited=deepcopy(existing);edited[0]['Title']='Human correction'
    assert action_reference(edited,envelope,expected_revision=snapshot['source_revision'],prior=reference['receipt'])['status']=='HUMAN_EDIT_OR_SOURCE_CHANGE'
    # The lifecycle starts from a separately supplied fictional state, never from an uncommitted proposal.
    state=lifecycle_state();binding=state['binding']
    binding.update(action_id=identity,register_revision=current['version'],minutes_record_id=s['source_ref'],minutes_revision=revision)
    state['current_binding']=deepcopy(binding)
    answered=response(action_id=identity)
    following=propose_lifecycle(state,req('record_response',state,response=answered))
    assert following['ok'];state=following['proposal']
    done=closure();done.update(decision_ref=s['source_ref'],decision_revision=revision)
    assert propose_lifecycle(state,req('resolve',state,closure=done))['proposal']['status']=='resolved'
    stale=deepcopy(state);stale['current_binding']['minutes_revision']='changed'
    assert propose_lifecycle(stale,req('resolve',stale,closure=done))['code']=='SOURCE_BINDING_CHANGED'
    cycle=dict(cycle='FICTIONAL-CYCLE-1',meeting=snapshot['meeting_id'],priorities='Review unanswered concern',
               pack='fixture://cycle/pack-v1',register='fixture://cycle/register',previous_cycle=None,
               controls={'papers':dict(state='Ready for chair review',note='Fictional papers reviewed',source='fixture://cycle/pack-v1'),
                         'concerns':dict(state='Waiting for input',note='Unresolved concern',source='fixture://cycle/concern')})
    nxt=next_cycle(cycle,'FICTIONAL-CYCLE-2','fixture://cycle/prior',control_keys=list(cycle['controls']))
    assert nxt['controls']['concerns']==cycle['controls']['concerns']
    assert nxt['controls']['papers']['state']=='Not started'
    feedback=summarise([dict(meeting_id=snapshot['meeting_id'],question_version='1',response_id='FICTIONAL-1',
                             recommendation='9',usefulness='5',comment='Private fictional comment')],snapshot['meeting_id'],'1')
    assert feedback['response_count']==1 and 'Private fictional comment' not in str(feedback)
    assert (cfg,data)==before


@pytest.mark.parametrize('fault',['source-version','meeting','paper-version','missing-paper-review','source-bytes'])
def test_changed_inputs_stop_the_affected_stage_without_output(tmp_path,fault):
    cfg,data=fixture(tmp_path)
    if fault=='source-version':
        cfg['rows'][0]['OData__UIVersionString']='2.0'
        with pytest.raises(ValueError):register_source_snapshot(**cfg)
    elif fault=='meeting':
        cfg['meeting']['revision']='event-changed'
        with pytest.raises(ValueError):register_source_snapshot(**cfg)
    elif fault=='source-bytes':
        Path(data['sources'][0]['path']).write_text('Changed fictional source')
        with pytest.raises(ValueError):minutes.build(data,tmp_path/'blocked.docx')
        assert not (tmp_path/'blocked.docx').exists()
    else:
        pack=json.loads((CASES/'pack-single.json').read_text())
        pack['papers'][0]['reviewed_version']='stale' if fault=='paper-version' else ''
        assert preflight(pack)['ready_for_manual_review'] is False


def test_ddms_metadata_recovery_and_register_recovery_preserve_independent_human_work():
    event,episode=contribution_fixture();original=deepcopy((event,episode))
    proposal=prepare_review(event,episode,event_version='1.0',target_etag='"1"',actual_work_confirmed=True,
        flags_confirmed=True,worker_confirmed=True,performer_reference='worker:opaque-123')
    saved={**episode,**proposal['patch'],'odata.etag':'"2"','OData__UIVersionString':'2.0'}
    assert reconcile_review(event,saved,proposal,performer_reference='worker:opaque-123')=='applied'
    assert saved['Coordination_x0020_note_x0020__x']==episode['Coordination_x0020_note_x0020__x']
    assert 'clinical_narrative' not in proposal['patch']
    cfg=recovery_config();plan=plan_action_register_recovery(**cfg)
    observed=dict(cfg['current'],**{cfg['field']:cfg['baseline'][cfg['field']],'odata.etag':'"4"'})
    assert reconcile_selective_recovery(plan,observed)=='RESTORED_READBACK'
    assert observed['HumanNote']==cfg['current']['HumanNote']
    assert (event,episode)==original
