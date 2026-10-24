import copy
import importlib.util
import json
from hashlib import sha256
from pathlib import Path
import pytest
import sys
from docx import Document

spec = importlib.util.spec_from_file_location('minutes', Path(__file__).parents[1] / 'scripts/prepare_minutes_draft.py')
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

@pytest.fixture
def data(tmp_path):
    source = tmp_path / 'fictional-register.txt'
    source.write_text('Fictional source for local software tests only.\nF-04:closure-accepted chair accepted closure for F-04.\nF-01-A:closure-accepted chair accepted closure for F-01-A.\n')
    row = dict(id='F-01', type='action', text='Review fictional process', status='Open', state='open', authority='Fictional approved instruction', owner='Fictional role, descriptive only', review_date=None, due_date=None, conditions='Only if capacity permits', dissent='Alternative remains unresolved', dependencies='Await fictional response', source_id='S1', verified=True, commissioning='commissioned', issue_readiness='hold', review_status='source_reviewed', closure_date=None, closure_authority=None, closure_acceptance='unknown', closure_evidence_ref=None)
    return dict(meeting_id='Fictional meeting', meeting_date='2026-09-28', cutoff='2026-09-14', scope='Fictional test only', sources=[dict(id='S1',path=str(source),sha256=sha256(source.read_bytes()).hexdigest(),link='https://example.invalid/register',revision='Fictional 1',snapshot='2026-09-14T09:00:00+10:00',access_scope='approved_for_local_review')], records=[row])

@pytest.fixture
def minutes_notes_data():
    sources = []
    for ident, fixture in m.F08_FIXTURES.items():
        sources.append(dict(id=ident, path=str(fixture['path']), sha256=fixture['sha256'],
                            link=fixture['link'], revision=fixture['revision'],
                            snapshot=fixture['snapshot'], access_scope='approved_for_local_review'))
    return dict(
        meeting_id=m.F08_MEETING_ID, meeting_date='2026-09-14', cutoff='2026-09-14',
        scope='Fictional software test only', sources=sources, records=[],
        minutes_notes=dict(schema_version=1, fixture_only=True,
                           agenda_source_id='F08-AGENDA-01', agenda_revision='Fictional agenda revision 1',
                           notes_source_id='F08-NOTES-01', notes_revision='Fictional notes revision 1',
                           entries=[
                               dict(agenda_item_ref='A-1', note_ref='N-1', kind='discussion'),
                               dict(agenda_item_ref='A-2', note_ref='N-2', kind='reported_decision'),
                               dict(agenda_item_ref='A-3', note_ref='N-3', kind='action_candidate'),
                               dict(agenda_item_ref=None, note_ref='N-ATT', kind='attendance'),
                               dict(agenda_item_ref=None, note_ref='N-4', kind='off_agenda_concern'),
                           ]))

@pytest.mark.parametrize('schema_version', [True, False, 1.0, '1'])
def test_f08_rejects_noninteger_schema_version_before_output(minutes_notes_data, tmp_path, schema_version):
    minutes_notes_data['minutes_notes']['schema_version'] = schema_version
    output = tmp_path / 'invalid-schema.docx'
    with pytest.raises(ValueError, match='F08 source identity or revision'):
        m.build(minutes_notes_data, output)
    assert not output.exists()

def test_f08_synthetic_notes_render_with_labels_unresolved_and_locators(minutes_notes_data, tmp_path):
    out=m.build(minutes_notes_data,tmp_path/'f08.docx')
    doc=Document(out)
    text='\n'.join(p.text for p in doc.paragraphs)
    assert doc.paragraphs[0].text == 'SYNTHETIC TEST OUTPUT — NOT ACTUAL MINUTES — NOT FOR CIRCULATION'
    banner=doc.paragraphs[0]
    assert banner.alignment == m.WD_ALIGN_PARAGRAPH.CENTER
    assert banner.runs[0].bold and banner.runs[0].font.size.pt == 14
    assert str(banner.runs[0].font.color.rgb) == '9C0006'
    assert banner._p.pPr.find(m.qn('w:shd')).get(m.qn('w:fill')) == 'FCE4D6'
    assert 'Agenda A-1: Review the fictional service dashboard and questions.' in text
    assert 'Discussion source note.' in text and 'N-1' in text
    assert 'Reported decision source note — authority unverified; not an authorised decision.' in text
    assert 'Action candidate source note — not commissioned, assigned or issued.' in text
    assert 'UNRESOLVED — No source-linked note supplied; outcome remains to be confirmed.' in text
    assert 'Off-agenda concerns — unresolved' in text
    assert 'existing concern route remains available without a paper or proposed solution' in text
    assert 'Urgent concerns are not delayed' in text
    assert 'No register write, assignment, due date, publication, notification or closure is inferred.' in text
    assert 'fixture://F08-DEMO-01/agenda/revision-1' in text
    assert 'F08-NOTES-01#F08-DEMO-01:N-2' in text

def test_f08_rejects_arbitrary_source_path_before_any_file_probe(minutes_notes_data, monkeypatch, tmp_path):
    minutes_notes_data['sources'][0]['path']=str(tmp_path/'private-notes.txt')
    def forbidden(*_args,**_kwargs):
        raise AssertionError('Unpinned source path must be rejected before filesystem access.')
    monkeypatch.setattr(Path,'is_file',forbidden)
    monkeypatch.setattr(Path,'is_symlink',forbidden)
    monkeypatch.setattr(Path,'read_bytes',forbidden)
    with pytest.raises(ValueError,match='pinned fixture manifest'):
        m.validate_minutes_notes(minutes_notes_data)

def test_f08_rejects_symlinked_fixture_before_read(minutes_notes_data, monkeypatch):
    original=m.Path.is_symlink
    def symlink_for_agenda(path):
        if path == m.F08_FIXTURES['F08-AGENDA-01']['path']:
            return True
        return original(path)
    def forbidden(*_args,**_kwargs):
        raise AssertionError('Pinned symlink must be rejected before file or content read.')
    monkeypatch.setattr(Path,'is_symlink',symlink_for_agenda)
    monkeypatch.setattr(Path,'is_file',forbidden)
    monkeypatch.setattr(Path,'read_bytes',forbidden)
    with pytest.raises(ValueError,match='must not be a symlink'):
        m.validate_minutes_notes(minutes_notes_data)

def test_f08_rejects_ancestor_symlink_before_file_probe(minutes_notes_data, monkeypatch):
    original=m.Path.resolve
    fixture=m.F08_FIXTURES['F08-AGENDA-01']['path']
    def alternate_resolve(path, *args, **kwargs):
        if path == fixture:
            return fixture.parent.parent/'symlinked-parent'/fixture.name
        return original(path,*args,**kwargs)
    def forbidden(*_args,**_kwargs):
        raise AssertionError('Ancestor symlink must fail before is_file/read_bytes.')
    monkeypatch.setattr(Path,'resolve',alternate_resolve)
    monkeypatch.setattr(Path,'is_file',forbidden)
    monkeypatch.setattr(Path,'read_bytes',forbidden)
    with pytest.raises(ValueError,match='not canonical'):
        m.validate_minutes_notes(minutes_notes_data)

def test_f08_rejects_output_outside_os_temp_before_writing(minutes_notes_data, tmp_path, monkeypatch):
    # A downloaded source archive may itself live under the OS temporary directory.
    # Control the boundary so the rejected path is outside it on every checkout.
    allowed=tmp_path/'f08-temp'; allowed.mkdir()
    monkeypatch.setattr(m.tempfile,'gettempdir',lambda: str(allowed))
    outside=tmp_path/'outside-f08.docx'
    with pytest.raises(ValueError,match='canonical OS temporary directory'):
        m.build(minutes_notes_data,outside)
    assert not outside.exists()

@pytest.mark.parametrize('mutation, message', [
    ('duplicate_note','unknown, duplicated'), ('missing_note','Every pinned notes passage'),
    ('bad_agenda_ref','existing agenda passage'), ('concern_item_link','must not be linked'),
    ('bad_hash','pinned fixture manifest'), ('extra_field','exact F08 minutes_notes fields'),
])
def test_f08_structural_mapping_failures_stop_before_output(minutes_notes_data,tmp_path,mutation,message):
    notes=minutes_notes_data['minutes_notes']
    if mutation=='duplicate_note': notes['entries'].append(copy.deepcopy(notes['entries'][0]))
    elif mutation=='missing_note': notes['entries'].pop()
    elif mutation=='bad_agenda_ref': notes['entries'][0]['agenda_item_ref']='A-UNKNOWN'
    elif mutation=='concern_item_link': notes['entries'][-1]['agenda_item_ref']='A-1'
    elif mutation=='bad_hash': minutes_notes_data['sources'][0]['sha256']='0'*64
    else: notes['unsafe_text']='Caller supplied text must not be accepted.'
    output=tmp_path/'blocked.docx'
    with pytest.raises(ValueError,match=message):m.build(minutes_notes_data,output)
    assert not output.exists()

def test_f08_repeat_does_not_overwrite_synthetic_output(minutes_notes_data,tmp_path):
    output=m.build(minutes_notes_data,tmp_path/'repeat.docx')
    before=output.read_bytes()
    with pytest.raises(FileExistsError):m.build(minutes_notes_data,output)
    assert output.read_bytes()==before

def test_selection_and_fidelity(data,tmp_path):
    data['records'][0].update(issue_readiness='issued', due_date='2026-09-28')
    for ident,state,typ,verified,dt in [('F-02','closure_pending','action',True,None),('F-03','follow_up','decision',True,'2026-09-20'),('F-04','closed','action',True,None),('F-05','candidate','action',False,None),('F-06','open','action',True,'2026-10-01'),('F-07','no_follow_up','decision',True,None)]:
        r=copy.deepcopy(data['records'][0]);r.update(id=ident,state=state,type=typ,verified=verified,review_date=dt)
        if ident=='F-06':r['due_date']='2026-10-01'
        if typ=='decision':r.update(commissioning='not_applicable',issue_readiness='not_applicable',review_status='approved_record')
        if state=='closed':r.update(closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S1#F-04:closure-accepted')
        data['records'].append(r)
    selected,excluded,_=m.select(data)
    assert [r['id'] for r in selected]==['F-01','F-02','F-03']
    assert [r['id'] for r,_ in excluded]==['F-04','F-05','F-06','F-07']
    out=m.build(data,tmp_path/'draft.docx');doc=Document(out);text='\n'.join(p.text for p in doc.paragraphs)
    for field in ['text','authority','conditions','dissent','dependencies','owner']:
        assert data['records'][0][field] in text
    assert 'Not recorded — review required' in text
    assert 'Closure remains awaiting acceptance' in text
    assert 'Attendance and new outcomes are not yet recorded' in text
    assert f"SHA-256: {data['sources'][0]['sha256']}" in text
    assert 'does not authenticate permission, source authority or factual completeness' in text
    assert 'excluded source IDs to reconcile before circulation' in text
    assert all(p.style.name=='Normal' for p in doc.paragraphs)
    assert 'F-05 | Source: S1 | Recorded snapshot status: Open | Unverified' in text
    assert 'F-04 | Source: S1 | Recorded snapshot status: Open | Recorded closed or no explicit follow-up | Closure reference: S1#F-04:closure-accepted' in text

@pytest.mark.parametrize('change',['hash','missing','duplicate','unknown_source','bad_date','cutoff','snapshot','type_state','string_verified','assignment','blank_condition'])
def test_fails_before_output(data,tmp_path,change):
    r=data['records'][0]
    if change=='hash':Path(data['sources'][0]['path']).write_text('Changed')
    elif change=='missing':data['sources'][0]['path']=str(tmp_path/'missing')
    elif change=='duplicate':data['records'].append(copy.deepcopy(r))
    elif change=='unknown_source':r['source_id']='Other'
    elif change=='bad_date':data['meeting_date']=None
    elif change=='cutoff':data['cutoff']='2026-10-01'
    elif change=='snapshot':data['sources'][0]['snapshot']='2026-09-15T09:00:00+10:00'
    elif change=='type_state':r['state']='follow_up'
    elif change=='string_verified':r['verified']='true'
    elif change=='assignment':r['assigned_to']='Person'
    else:r['conditions']=' '
    out=tmp_path/'draft.docx'
    with pytest.raises((ValueError,TypeError)):m.build(data,out)
    assert not out.exists()

def test_distinct_subjects_reusing_legacy_id_fail_closed(data):
    conflicting = copy.deepcopy(data['records'][0])
    conflicting.update(text='Ethics and Clinical Trials / QHEPS')
    data['records'][0].update(text='Dictation system / Winscribe')
    data['records'].append(conflicting)
    with pytest.raises(ValueError, match='Duplicate record ID; reconcile sources before drafting'):
        m.select(data)

def test_draft_only_action_omitted_from_later_source_is_not_carried_forward(data, tmp_path):
    later = tmp_path / 'later-minutes.txt'
    later.write_text('Fictional later minutes; no such action recorded.\n')
    data['sources'].append(dict(id='S2', path=str(later), sha256=sha256(later.read_bytes()).hexdigest(),
                                link='https://example.invalid/later', revision='Fictional approved candidate',
                                snapshot='2026-09-14T10:00:00+10:00', access_scope='approved_for_local_review'))
    data['records'][0].update(source_id='S1', review_status='draft', verified=True,
                              issue_readiness='hold', text='Fictional draft-only instruction')
    selected, excluded, _ = m.select(data)
    assert selected == []
    assert excluded[0][1] == 'Action source remains draft; excluded from action carry-forward'

def test_date_only_confirmation_does_not_approve_decision_file(data):
    data['records'][0].update(type='decision', state='follow_up', commissioning='not_applicable',
                              issue_readiness='not_applicable', review_status='source_reviewed',
                              authority='Fictional date confirmed; prior file identity not recorded')
    selected, excluded, _ = m.select(data)
    assert selected == []
    assert excluded[0][1] == 'Decision source not approved; excluded from decision carry-forward'

@pytest.mark.parametrize('status', ['All previous action items addressed', 'CLOSE ITEM; date conflicts with later updates'])
def test_collective_or_conflicting_closure_cue_remains_pending(data, tmp_path, status):
    source = Path(data['sources'][0]['path'])
    source.write_text('Fictional collective cue; no item-specific accepted closure.\n')
    data['sources'][0]['sha256'] = sha256(source.read_bytes()).hexdigest()
    data['records'][0].update(state='closure_pending', status=status,
                              closure_acceptance='reported_complete', closure_date=None,
                              closure_authority=None, closure_evidence_ref=None,
                              issue_readiness='issued', due_date='2026-09-28')
    selected, excluded, _ = m.select(data)
    assert len(selected) == 1 and not excluded
    out = m.build(data, tmp_path / 'closure-pending.docx')
    text = '\n'.join(p.text for p in Document(out).paragraphs)
    assert 'Closure remains awaiting acceptance; reported completion is not accepted closure.' in text
    assert f'Recorded status: {status}' in text

@pytest.mark.parametrize('scope', ['restricted_pending', 'unrecognised_scope'])
def test_source_access_scope_fails_before_opening_path(data, tmp_path, monkeypatch, scope):
    data['sources'][0].update(path=str(tmp_path / 'must-not-open.txt'), access_scope=scope)
    def unexpected_open(_self):
        raise AssertionError('Source path was inspected before access scope was checked.')
    monkeypatch.setattr(Path, 'is_file', unexpected_open)
    expected = 'Restricted source is not read' if scope == 'restricted_pending' else 'explicit known access-scope'
    with pytest.raises(ValueError, match=expected):
        m.select(data)

@pytest.mark.parametrize('failure', ['existing_output', 'output_access', 'invalid_input'])
def test_cli_reports_safe_actionable_failure(data, tmp_path, monkeypatch, capsys, failure):
    if failure == 'invalid_input':
        data.pop('scope')
    input_path = tmp_path / 'fictional-input.json'
    input_path.write_text(json.dumps(data), encoding='utf-8')
    output = tmp_path / 'missing' / 'draft.docx' if failure == 'output_access' else tmp_path / 'draft.docx'
    if failure == 'existing_output':
        output.write_bytes(b'Fictional human edit; retain unchanged.')
    monkeypatch.setattr(sys, 'argv', ['prepare_minutes_draft.py', str(input_path), str(output)])
    with pytest.raises(SystemExit) as exc:
        m.main()
    assert exc.value.code == 1
    message = capsys.readouterr().err
    assert 'Draft not prepared' in message
    if failure == 'existing_output':
        assert 'Existing draft retained' in message
        assert output.read_bytes() == b'Fictional human edit; retain unchanged.'
    elif failure == 'output_access':
        assert 'file-access error' in message
    else:
        assert 'Use the documented meeting, scope, source and record fields' in message
        assert not output.exists()
    assert str(input_path) not in message and str(output) not in message

def test_repeat_preserves_human_edits_and_original_source(data,tmp_path):
    source=Path(data['sources'][0]['path']);before=source.read_bytes()
    out=m.build(data,tmp_path/'draft.docx');doc=Document(out);doc.add_paragraph('Human correction');doc.save(out);edited=out.read_bytes()
    with pytest.raises(FileExistsError):m.build(data,out)
    assert out.read_bytes()==edited and source.read_bytes()==before

def test_empty_snapshot_not_claim_of_no_work(data,tmp_path):
    data['records']=[]
    out=m.build(data,tmp_path/'draft.docx')
    assert 'does not establish that no outstanding business exists' in '\n'.join(p.text for p in Document(out).paragraphs)


def test_docx_closure_waits_for_revision_aware_source_review(data, tmp_path):
    source = tmp_path / 'fictional-register.docx'
    doc = Document()
    doc.add_paragraph('F-04:closure-accepted chair accepted closure for F-04.')
    doc.save(source)
    data['sources'][0].update(path=str(source), sha256=sha256(source.read_bytes()).hexdigest())
    data['records'][0].update(state='closed', closure_date='2026-09-10',
                              closure_authority='Fictional chair accepted closure',
                              closure_acceptance='accepted',
                              closure_evidence_ref='S1#F-04:closure-accepted')
    with pytest.raises(ValueError, match='Accepted closure needs'):
        m.select(data)


def test_duplicate_closure_locator_requires_reconciliation(data):
    source = Path(data['sources'][0]['path'])
    source.write_text('F-01:closure-review chair accepted closure.\nF-01:closure-review later note disputes closure.\n')
    data['sources'][0]['sha256'] = sha256(source.read_bytes()).hexdigest()
    data['records'][0].update(state='closed', closure_date='2026-09-10',
                              closure_authority='Fictional chair accepted closure',
                              closure_acceptance='accepted',
                              closure_evidence_ref='S1#F-01:closure-review')
    with pytest.raises(ValueError, match='Accepted closure needs'):
        m.select(data)


def test_closure_locator_prefix_of_another_passage_key_is_rejected(data):
    source = Path(data['sources'][0]['path'])
    source.write_text('F-01:closure-revoked chair did not accept closure.\n')
    data['sources'][0]['sha256'] = sha256(source.read_bytes()).hexdigest()
    data['records'][0].update(state='closed', closure_date='2026-09-10',
                              closure_authority='Fictional chair accepted closure',
                              closure_acceptance='accepted',
                              closure_evidence_ref='S1#F-01:closure')
    with pytest.raises(ValueError, match='Accepted closure needs'):
        m.select(data)


def test_snapshot_cutoff_uses_brisbane_date(data):
    # This is 15 September in Brisbane, despite the timestamp's 14 September date.
    data['sources'][0]['snapshot']='2026-09-14T23:30:00+00:00'
    with pytest.raises(ValueError, match='must not follow cut-off'):
        m.select(data)

    # This is still 14 September in Brisbane, despite its 15 September date.
    data['sources'][0]['snapshot']='2026-09-15T00:30:00+12:00'
    selected, excluded, _ = m.select(data)
    assert not selected and len(excluded) == 1


def test_overdue_date_preserved_despite_later_review(data,tmp_path):
    data['records'][0].update(due_date='2026-09-01',review_date='2026-10-15',issue_readiness='issued')
    selected,excluded,_=m.select(data)
    assert len(selected)==1 and not excluded
    out=m.build(data,tmp_path/'overdue.docx')
    text='\n'.join(p.text for p in Document(out).paragraphs)
    assert 'Recorded due date: 2026-09-01' in text
    assert 'Review eligibility: 2026-10-15' in text
    assert 'Due/review:' not in text


@pytest.mark.parametrize('commissioning',['unclear','not_commissioned'])
def test_source_verified_action_without_commission_is_excluded(data,commissioning):
    data['records'][0]['commissioning']=commissioning
    selected,excluded,_=m.select(data)
    assert selected==[]
    assert excluded[0][1]=='Commissioning not evidenced; excluded from action carry-forward'


@pytest.mark.parametrize('readiness', ['hold', 'ready'])
def test_commissioned_but_not_issued_stays_in_preparation_review(data,readiness):
    data['records'][0]['issue_readiness']=readiness
    if readiness=='ready':data['records'][0]['review_date']='2026-09-20'
    selected,excluded,_=m.select(data)
    assert selected==[]
    expected='Action not issued; retain for preparation review only'
    if readiness=='hold':
        expected+=' — agreed due/review date not recorded in supplied snapshot'
    assert excluded[0][1]==expected


@pytest.mark.parametrize('owner,date_fields,expected', [
    ('Not recorded', (None, None), ('accountable owner not evidenced/accepted in supplied snapshot', 'agreed due/review date not recorded in supplied snapshot')),
    ('Fictional role, descriptive only', (None, None), ('', 'agreed due/review date not recorded in supplied snapshot')),
    ('To confirm', ('2026-09-20', None), ('accountable owner not evidenced/accepted in supplied snapshot', '')),
    ('Not recorded', (None, '2026-09-20'), ('accountable owner not evidenced/accepted in supplied snapshot', '')),
])
def test_not_issued_exclusion_exposes_all_missing_particulars_without_mutation(data,owner,date_fields,expected):
    row=data['records'][0]
    row.update(issue_readiness='hold', owner=owner, due_date=date_fields[0], review_date=date_fields[1])
    original=json.loads(json.dumps(data))
    selected,excluded,_=m.select(data)
    assert selected==[]
    reason=excluded[0][1]
    assert reason.startswith('Action not issued; retain for preparation review only')
    for phrase in expected:
        if phrase:
            assert phrase in reason
    assert data==original
    assert excluded[0][0]['source_id']=='S1'
    assert excluded[0][0]['status']==original['records'][0]['status']


def test_draft_action_cannot_carry_forward_despite_verified_flag(data):
    data['records'][0]['review_status']='draft'
    selected,excluded,_=m.select(data)
    assert selected==[]
    assert excluded[0][1]=='Action source remains draft; excluded from action carry-forward'


@pytest.mark.parametrize('review_status',['draft','source_reviewed'])
def test_provisional_decision_is_excluded(data,review_status):
    data['records'][0].update(type='decision',state='follow_up',commissioning='not_applicable',issue_readiness='not_applicable',review_status=review_status)
    selected,excluded,_=m.select(data)
    assert selected==[]
    assert excluded[0][1]=='Decision source not approved; excluded from decision carry-forward'


@pytest.mark.parametrize('change',['uncommissioned_ready','undated_ready','draft_ready','missing_issue_authority','placeholder_issue_owner','placeholder_issue_authority','to_confirm_closure_authority','unaccepted_closure','missing_closure_authority','placeholder_closure_authority','reported_complete_as_closure','reported_completion_variant','reported_complete_open_state','missing_closure_reference','unknown_closure_source','blank_closure_locator','absent_closure_locator','other_row_closure_locator','prefix_collision_locator','closed_status_open_state','missing_decision_authority','to_confirm_decision_authority'])
def test_invalid_issue_or_closure_fails(data,change):
    row=data['records'][0]
    if change=='uncommissioned_ready':row.update(commissioning='unclear',issue_readiness='ready',review_date='2026-09-20')
    elif change=='undated_ready':row['issue_readiness']='ready'
    elif change=='draft_ready':row.update(issue_readiness='ready',review_date='2026-09-20',review_status='draft')
    elif change=='missing_issue_authority':row.update(issue_readiness='ready',review_date='2026-09-20',authority='Not recorded')
    elif change=='placeholder_issue_owner':row.update(issue_readiness='ready',review_date='2026-09-20',owner='To confirm; no person assignment')
    elif change=='placeholder_issue_authority':row.update(issue_readiness='ready',review_date='2026-09-20',authority='Unverified historical source')
    elif change=='unaccepted_closure':row['state']='closed'
    elif change=='missing_closure_authority':row.update(state='closed',closure_date='2026-09-10',closure_acceptance='accepted',closure_evidence_ref='S1#fictional-closure')
    elif change=='placeholder_closure_authority':row.update(state='closed',closure_date='2026-09-10',closure_authority='Not recorded',closure_acceptance='accepted',closure_evidence_ref='S1#fictional-closure')
    elif change=='to_confirm_closure_authority':row.update(state='closed',closure_date='2026-09-10',closure_authority='To confirm',closure_acceptance='accepted',closure_evidence_ref='S1#fictional-closure')
    elif change=='reported_complete_as_closure':row.update(state='closed',status='Reported complete',closure_date='2026-09-10',closure_authority='Reported complete',closure_acceptance='reported_complete',closure_evidence_ref='S1#fictional-report')
    elif change=='reported_completion_variant':row.update(state='closed',status='Reported complete',closure_date='2026-09-10',closure_authority='Reported completion confirmed by service',closure_acceptance='reported_complete',closure_evidence_ref='S1#fictional-report')
    elif change=='reported_complete_open_state':row.update(state='open',status='Reported complete',closure_acceptance='reported_complete',issue_readiness='issued',due_date='2026-09-01')
    elif change=='missing_closure_reference':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted')
    elif change=='unknown_closure_source':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S999#fictional-closure')
    elif change=='blank_closure_locator':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S1#')
    elif change=='absent_closure_locator':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S1#LOCATOR-DOES-NOT-EXIST-9F31')
    elif change=='other_row_closure_locator':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S1#F-04:closure-accepted')
    elif change=='prefix_collision_locator':row.update(state='closed',closure_date='2026-09-10',closure_authority='Fictional chair accepted closure',closure_acceptance='accepted',closure_evidence_ref='S1#F-01-A:closure-accepted')
    elif change=='closed_status_open_state':row.update(status='Closed — awaiting reconciliation',state='open',issue_readiness='issued',due_date='2026-09-01')
    elif change=='to_confirm_decision_authority':row.update(type='decision',state='follow_up',commissioning='not_applicable',issue_readiness='not_applicable',review_status='approved_record',authority='To confirm')
    else:row.update(type='decision',state='follow_up',commissioning='not_applicable',issue_readiness='not_applicable',review_status='approved_record',authority='Not recorded')
    with pytest.raises(ValueError):m.select(data)
