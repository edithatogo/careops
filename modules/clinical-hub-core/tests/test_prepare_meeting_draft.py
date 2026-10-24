import copy
import importlib.util
import json
from io import BytesIO
from pathlib import Path

from docx import Document
import pytest

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('prepare_meeting_draft', ROOT / 'scripts/prepare_meeting_draft.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


@pytest.fixture
def data():
    return json.loads((ROOT / 'tests/fixtures/reviewed_agenda.json').read_text())


def test_memory_builder_preserves_content_links_and_input_without_writes(data, tmp_path, monkeypatch):
    data['items'][0]['source'] = 'https://example.test/reviewed.pdf'
    original = copy.deepcopy(data)
    monkeypatch.chdir(tmp_path)
    import builtins
    import io
    real_open = io.open

    def read_only_open(file, mode='r', *args, **kwargs):
        if any(flag in mode for flag in 'wax+'):
            raise AssertionError('In-memory builder attempted a file write')
        return real_open(file, mode, *args, **kwargs)

    monkeypatch.setattr(builtins, 'open', read_only_open)
    monkeypatch.setattr(io, 'open', read_only_open)
    blob = module.build_bytes(data)
    assert isinstance(blob, bytes)
    assert blob == module.build_bytes(data)
    doc = Document(BytesIO(blob))
    paragraphs = [p.text for p in doc.paragraphs]
    for item in data['items']:
        assert item['body'] in paragraphs
    links = [rel.target_ref for rel in doc.part.rels.values()
             if rel.reltype.endswith('/hyperlink')]
    assert data['items'][0]['source'] in links
    assert data == original
    assert list(tmp_path.iterdir()) == []


def test_file_wrapper_matches_memory_bytes_and_validates_once(data, tmp_path, monkeypatch):
    expected = module.build_bytes(data)
    calls = []
    original_validate = module.validate

    def counted_validate(value):
        calls.append(value)
        return original_validate(value)

    monkeypatch.setattr(module, 'validate', counted_validate)
    result = module.build(data, tmp_path / 'shared.docx')
    assert result.read_bytes() == expected
    assert len(calls) == 1


def test_memory_builder_rejects_unreviewed_input(data):
    data['items'][0]['reviewed'] = False
    with pytest.raises(ValueError, match='explicit human review'):
        module.build_bytes(data)


def test_preserves_reviewed_content_and_standing_order(data, tmp_path):
    result = module.build(data, tmp_path / 'draft.docx')
    doc = Document(result)
    paragraphs = [p.text for p in doc.paragraphs]
    text = '\n'.join(paragraphs)
    assert 'UNRESOLVED' in text
    assert text.index('Guests — early contributions') < text.index('Update from the Executive') < text.index('Matters for noting')
    for item in data['items']:
        assert item['body'] in paragraphs
        assert f"Source: {item['source']}" in paragraphs
        assert f"Purpose: {item['purpose']}" in paragraphs
    assert all(p.style.name == 'Normal' for p in doc.paragraphs)
    assert all(run._element.rPr is None for p in doc.paragraphs for run in p.runs)


@pytest.mark.parametrize('source', [
    'https://example.test/papers/reviewed.pdf?version=3',
    'https://example.test:8443/papers/reviewed.pdf?version=3',
])
def test_https_source_is_clickable_and_source_text_is_preserved(data, tmp_path, source):
    data['items'][0]['source'] = source
    doc = Document(module.build(data, tmp_path / 'linked-source.docx'))
    source_paragraph = next(p for p in doc.paragraphs if p.text.startswith('Source: '))
    paragraphs = [p.text for p in doc.paragraphs]
    assert paragraphs.index(source_paragraph.text) > paragraphs.index(data['items'][0]['body'])
    assert paragraphs.index(source_paragraph.text) < next(
        index for index, text in enumerate(paragraphs)
        if text.startswith('Existing action reference:')
    )
    hyperlinks = source_paragraph._p.xpath('./w:hyperlink')
    assert len(hyperlinks) == 1
    relation_id = hyperlinks[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    relation = doc.part.rels[relation_id]
    assert relation.reltype.endswith('/hyperlink')
    assert relation.is_external
    assert relation.target_ref == source
    visible_link_text = ''.join(hyperlinks[0].xpath('.//w:t/text()'))
    assert visible_link_text == source


@pytest.mark.parametrize(('source', 'message'), [
    ('https:///papers/reviewed.pdf', 'needs a hostname'),
    ('https://exa mple.test/papers/reviewed.pdf', 'must not contain whitespace'),
    ('https://example.test/papers/reviewed file.pdf', 'must not contain whitespace'),
    ('https://example.test/papers/reviewed.pdf?label=reviewed source', 'must not contain whitespace'),
    ('https://example.test\\evil/papers/reviewed.pdf', 'must not contain a backslash'),
    ('https://example.test:/papers/reviewed.pdf', 'empty port'),
    ('https://example.test:notaport/papers/reviewed.pdf', 'invalid port'),
    ('https://example.test:65536/papers/reviewed.pdf', 'invalid port'),
])
def test_invalid_https_source_is_rejected_before_output(data, tmp_path, source, message):
    data['items'][0]['source'] = source
    output = tmp_path / 'invalid-source.docx'
    with pytest.raises(ValueError, match=message):
        module.build(data, output)
    assert not output.exists()


def test_non_url_source_identifier_remains_plain_visible_text(data, tmp_path):
    source = 'synthetic:SYN-01'
    data['items'][0]['source'] = source
    doc = Document(module.build(data, tmp_path / 'opaque-source.docx'))
    source_paragraph = next(p for p in doc.paragraphs if p.text.startswith('Source: '))
    assert source_paragraph.text == f'Source: {source}'
    assert not source_paragraph._p.xpath('./w:hyperlink')


@pytest.mark.parametrize('mutation', ['duplicate', 'duplicate_content', 'source', 'credential_source', 'unreviewed', 'action', 'purpose', 'date'])
def test_rejects_unsafe_input_before_output(data, tmp_path, mutation):
    if mutation == 'duplicate':
        data['items'].append(copy.deepcopy(data['items'][0]))
    elif mutation == 'duplicate_content':
        item = copy.deepcopy(data['items'][0])
        item['id'] = 'NEW-ID'
        data['items'].append(item)
    elif mutation == 'source':
        data['items'][0]['source'] = ' '
    elif mutation == 'credential_source':
        data['items'][0]['source'] = 'https://synthetic-user:synthetic-password@example.test/paper.pdf'
    elif mutation == 'unreviewed':
        data['items'][0]['reviewed'] = False
    elif mutation == 'action':
        data['items'][0]['assigned_to'] = 'Fictional owner'
    elif mutation == 'purpose':
        data['items'][0]['purpose'] = 'assigned action'
    else:
        data['meeting_date'] = 'tomorrow'
    output = tmp_path / 'draft.docx'
    with pytest.raises(ValueError):
        module.build(data, output)
    assert not output.exists()


def test_never_overwrites(data, tmp_path):
    output = tmp_path / 'draft.docx'
    output.write_bytes(b'Original must survive')
    with pytest.raises(FileExistsError):
        module.build(data, output)
    assert output.read_bytes() == b'Original must survive'


def test_reproducible_and_missing_sections_explicit(tmp_path):
    data = {'meeting_date': None, 'items': []}
    first = module.build(data, tmp_path / 'a.docx')
    second = module.build(data, tmp_path / 'b.docx')
    assert first.read_bytes() == second.read_bytes()
    text = '\n'.join(p.text for p in Document(first).paragraphs)
    assert 'Update from the Executive' in text
    assert text.count('Content unresolved') == 3


def test_rejects_text_output(data, tmp_path):
    with pytest.raises(ValueError):
        module.build(data, tmp_path / 'draft.txt')


def test_follow_through_survives_decision_focused_generation(data, tmp_path):
    data['items'][0]['purpose'] = 'decision'
    path = module.build(data, tmp_path / 'cycle.docx')
    text = '\n'.join(p.text for p in Document(path).paragraphs)
    assert text.index('Update from the Executive') < text.index('Previous minutes and corrections') < text.index('Matters for noting, discussion and decision')
    assert 'Responses due and matters held over' in text
    assert 'Completion is not accepted closure' in text
    assert 'Purpose and reviewed flags do not establish authority' in text
    assert 'do not copy progress history' in text
    assert 'read back decisions/recommendations, agreed action references and responses due' in text


def test_rejects_whitespace_only_retry_with_new_id(data, tmp_path):
    duplicate = copy.deepcopy(data['items'][0])
    duplicate['id'] = 'RETRY-TRANSPORT-ID'
    for key in ('title', 'body', 'source'):
        duplicate[key] = '  '.join(duplicate[key].split())
    data['items'].append(duplicate)
    output = tmp_path / 'retry.docx'
    with pytest.raises(ValueError, match='Duplicate item content'):
        module.build(data, output)
    assert not output.exists()


def test_rejects_duplicate_id_with_equivalent_internal_whitespace(data, tmp_path):
    data['items'][0]['id'] = 'CASE  01'
    duplicate = copy.deepcopy(data['items'][1])
    duplicate['id'] = ' case\t01 '
    duplicate['body'] = 'SYNTHETIC ONLY: a distinct agenda item with the same logical ID.'
    data['items'].append(duplicate)
    output = tmp_path / 'duplicate-id.docx'
    with pytest.raises(ValueError, match='Duplicate item id'):
        module.build(data, output)
    assert not output.exists()
