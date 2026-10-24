#!/usr/bin/env python3
"""Prepare a local review-only agenda from human-reviewed items; no delivery API."""
import argparse
from datetime import date, datetime, timezone
from io import BytesIO
import json
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt

PURPOSES = ('noting', 'discussion', 'decision')
SECTIONS = ('guests', 'executive', 'main')


def validate(data):
    if not isinstance(data, dict) or set(data) != {'meeting_date', 'items'}:
        raise ValueError('Input requires exactly meeting_date and items; action/assignment fields are not supported.')
    meeting_date = data['meeting_date']
    if meeting_date is not None:
        if not isinstance(meeting_date, str) or date.fromisoformat(meeting_date).isoformat() != meeting_date:
            raise ValueError('meeting_date must be ISO YYYY-MM-DD or null.')
    if not isinstance(data['items'], list):
        raise ValueError('items must be a list.')
    seen = set()
    seen_content = set()
    for item in data['items']:
        if not isinstance(item, dict) or set(item) != {'id', 'title', 'body', 'source', 'reviewed', 'purpose', 'section'}:
            raise ValueError('Each item requires id, title, body, source, reviewed, purpose and section only.')
        for key in ('id', 'title', 'body', 'source'):
            if not isinstance(item[key], str) or not item[key].strip():
                raise ValueError(f'Item {key} must contain text.')
            if any(ord(c) < 32 and c not in '\n\t\r' for c in item[key]):
                raise ValueError(f'Item {key} contains unsupported control characters.')
        try:
            parsed_source = urlsplit(item['source'])
        except ValueError as exc:
            raise ValueError('Source URL is malformed.') from exc
        if parsed_source.username is not None or parsed_source.password is not None:
            raise ValueError('Source URL must not contain embedded credentials.')
        if parsed_source.scheme.casefold() == 'https':
            if any(character.isspace() for character in item['source']):
                raise ValueError('HTTPS source URL must not contain whitespace.')
            if '\\' in item['source']:
                raise ValueError('HTTPS source URL must not contain a backslash.')
            if not parsed_source.hostname:
                raise ValueError('HTTPS source URL needs a hostname.')
            if parsed_source.netloc.endswith(':'):
                raise ValueError('HTTPS source URL has an empty port.')
            try:
                parsed_source.port
            except ValueError as exc:
                raise ValueError('HTTPS source URL has an invalid port.') from exc
        identity = ' '.join(item['id'].split()).casefold()
        if identity in seen:
            raise ValueError('Duplicate item id.')
        seen.add(identity)
        if item['reviewed'] is not True:
            raise ValueError('Every item needs explicit human review (reviewed: true).')
        if item['purpose'] not in PURPOSES or item['section'] not in SECTIONS:
            raise ValueError('Unsupported item purpose or section.')
        signature = tuple(' '.join(item[key].split()).casefold() for key in ('title', 'body', 'source'))
        if signature in seen_content:
            raise ValueError('Duplicate item content, even with a different id.')
        seen_content.add(signature)
    return data


def add_source(paragraph, source):
    """Keep source text visible and link only well-formed HTTPS URLs."""
    paragraph.add_run('Source: ')
    parsed = urlsplit(source)
    if (parsed.scheme.lower() == 'https' and parsed.hostname
            and not any(character.isspace() for character in source)):
        hyperlink = OxmlElement('w:hyperlink')
        relationship_id = paragraph.part.relate_to(source, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn('r:id'), relationship_id)
        hyperlink.append(paragraph.add_run(source)._r)
        paragraph._p.append(hyperlink)
    else:
        paragraph.add_run(source)


def build_bytes(data):
    """Return reproducible review-only DOCX bytes without writing an output file."""
    validate(data)
    document = Document()
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25.4)
    lines = [
        'Clinical Directors — draft agenda',
        'LOCAL REVIEW DRAFT — not issued; no decisions or actions have been recorded.',
        'Meeting date: ' + (data['meeting_date'] or 'UNRESOLVED — confirm before issue'),
        'Principles: respectful, candid discussion; restorative, just and learning culture; support departmental leadership.',
        'Welcome, apologies and conflicts of interest',
    ]
    for slot, title in [('guests', 'Guests — early contributions'), ('executive', 'Update from the Executive'), ('main', 'Matters for noting, discussion and decision')]:
        if slot == 'main':
            lines.extend(['', 'Previous minutes and corrections — confirm the reviewed version and preserve the original when correcting it.'])
        lines.extend(['', title])
        items = [item for item in data['items'] if item['section'] == slot]
        if not items:
            lines.append('Content unresolved — confirm before issue.')
        for item in items:
            lines.extend([f"{item['id']}: {item['title']}", f"Purpose: {item['purpose']}", item['body']])
            lines.append(('source', item['source']))
            lines.append('Existing action reference: confirm ID/link or none; do not copy progress history.')
            if item['purpose'] == 'decision':
                lines.append('Authority checkpoint: verify decision-maker and basis before issue; if unresolved, seek advice or a recommendation only. Purpose and reviewed flags do not establish authority.')
                lines.append('Confirm options, outcome sought and affected perspectives in the reviewed paper.')
    lines.extend(['', 'Existing action register — review agreed actions using the existing register; link and maintainer require confirmation if not already verified.',
                  'Responses due and matters held over — retain unanswered questions, reasons, agreed response contact and next review date. Completion is not accepted closure.',
                  'Remaining noting items — noting is not endorsement; raise material questions before closing.',
                  'Close — read back decisions/recommendations, agreed action references and responses due; confirm permitted departmental messages and next meeting arrangements.'])
    for line in lines:
        if isinstance(line, tuple) and line[0] == 'source':
            add_source(document.add_paragraph(style='Normal'), line[1])
        else:
            document.add_paragraph(line, style='Normal')
    properties = document.core_properties
    properties.author = properties.last_modified_by = properties.comments = ''
    properties.title = 'Clinical Directors draft agenda'
    properties.subject = 'Local review draft; not adopted or issued'
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    raw = BytesIO()
    document.save(raw)
    canonical = BytesIO()
    with ZipFile(raw) as source, ZipFile(canonical, 'w', compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            target.writestr(info, source.read(name))
    return canonical.getvalue()


def build(data, output):
    """Write the shared DOCX bytes to a new file; never overwrite an existing file."""
    output = Path(output)
    if output.suffix.lower() != '.docx':
        raise ValueError('Output must be a .docx file.')
    content = build_bytes(data)
    with output.open('xb') as handle:
        handle.write(content)
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, help='Local JSON with meeting_date and reviewed items')
    parser.add_argument('output', type=Path, help='New .docx path; existing files are never overwritten')
    args = parser.parse_args()
    try:
        build(json.loads(args.input.read_text()), args.output)
    except (ValueError, OSError) as error:
        parser.exit(1, f'Draft not prepared: {error}\n')
    print(f'Local review draft prepared: {args.output}')


if __name__ == '__main__':
    main()
