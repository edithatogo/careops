#!/usr/bin/env python3
"""Build one local minutes draft from reviewed register snapshots; no tenant access.

The JSON is a technical adapter interface, not a member submission form.
A source hash checks file integrity, not commissioning or factual correctness.
The source access declaration is caller-supplied and does not prove permission.
Existing outputs are never overwritten; source refresh needs human reconciliation.
Selection state closed requires separate accepted closure, dated authority and
a source locator; review_date is the
human-agreed review eligibility date, not a rescheduled due date. Source review
and commissioning are separate inputs, so a checked source alone cannot issue
work or promote a provisional decision. Include unresolved overdue work.
"""
import argparse
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

STATES = {'open', 'closure_pending', 'closed', 'candidate', 'follow_up', 'no_follow_up'}
ROW_FIELDS = {'id', 'type', 'text', 'status', 'state', 'authority', 'owner',
              'review_date', 'due_date', 'conditions', 'dissent', 'dependencies', 'source_id', 'verified',
              'commissioning', 'issue_readiness', 'review_status', 'closure_date', 'closure_authority',
              'closure_acceptance', 'closure_evidence_ref'}
UNRESOLVED_PREFIXES = ('not recorded', 'to confirm', 'tbc', 'unknown', 'unverified', 'pending')
MEETING_TIMEZONE = ZoneInfo('Australia/Brisbane')
F08_MEETING_ID = 'F08-DEMO-01'
F08_FIXTURES = {
    'F08-AGENDA-01': {
        'path': Path(__file__).resolve().parents[1] / 'tests/fixtures/f08/agenda.txt',
        'revision': 'Fictional agenda revision 1',
        'sha256': '01061b7d4425a5e9a9a019487e1ce1f0b58ee5f12f10746fc9c6e10f18bce4eb',
        'link': 'fixture://F08-DEMO-01/agenda/revision-1',
        'snapshot': '2026-09-14T09:00:00+10:00',
    },
    'F08-NOTES-01': {
        'path': Path(__file__).resolve().parents[1] / 'tests/fixtures/f08/notes.txt',
        'revision': 'Fictional notes revision 1',
        'sha256': '3e2b2fdf95b5cf405835a2cd432782d1644720e8cbf3d553e98fb16f0fbe5e7c',
        'link': 'fixture://F08-DEMO-01/notes/revision-1',
        'snapshot': '2026-09-14T09:05:00+10:00',
    },
}
F08_NOTE_KINDS = {
    'attendance', 'discussion', 'reported_decision', 'action_candidate',
    'condition', 'dissent', 'held', 'unresolved', 'off_agenda_concern',
}


def required_text(value):
    return isinstance(value, str) and bool(value.strip()) and not any(
        ord(c) < 32 and c not in '\n\r\t' for c in value)


def evidenced_text(value):
    """A non-empty placeholder is not evidence of an owner or authority."""
    if not required_text(value):
        return False
    normalized = ' '.join(value.casefold().split()).lstrip(' :;-')
    return normalized not in {'n/a', 'na', 'none'} and not normalized.startswith(UNRESOLVED_PREFIXES)


def known_source_locator(value, sources, record_id):
    """Check a record-bound literal locator in a hashed, readable text snapshot."""
    if not required_text(value) or '#' not in value:
        return False
    source_id, locator = value.split('#', 1)
    record_key, separator, passage_key = locator.partition(':')
    if (source_id not in sources or not separator or record_key != record_id
            or not evidenced_text(passage_key)):
        return False
    path = Path(sources[source_id]['path'])
    if path.suffix.lower() not in {'.txt', '.md', '.csv'}:
        return False
    try:
        matches = [
            line.strip().split(maxsplit=1)[0]
            for line in path.read_text(encoding='utf-8').splitlines()
            if line.strip()
        ]
        return matches.count(locator) == 1
    except UnicodeError:
        return False


def claims_closed(status):
    normalized = ' '.join(status.casefold().split()).lstrip(' :;-')
    return normalized == 'closed' or normalized.startswith(('closed ', 'closed:', 'accepted closure'))


def iso_date(value):
    if not isinstance(value, str) or date.fromisoformat(value).isoformat() != value:
        raise ValueError('Use an explicit ISO date.')
    return date.fromisoformat(value)


def _f08_lines(source):
    """Read one already-pinned, non-symlink fixture and parse opaque refs."""
    path = F08_FIXTURES[source['id']]['path']
    if path.resolve() != path:
        raise ValueError('Pinned fixture path is not canonical; reject symlink substitution.')
    if path.is_symlink():
        raise ValueError('Pinned fixture must not be a symlink.')
    if not path.is_file():
        raise ValueError('Pinned fixture path changed; reconcile before drafting.')
    raw = path.read_bytes()
    if sha256(raw).hexdigest() != source['sha256']:
        raise ValueError('Pinned fixture changed; reconcile before drafting.')
    try:
        lines = raw.decode('utf-8').splitlines()
    except UnicodeError as error:
        raise ValueError('Pinned fixture is not UTF-8 text.') from error
    parsed = {}
    for line in lines:
        if not line:
            continue
        left, separator, literal = line.partition(' | ')
        meeting, colon, ref = left.partition(':')
        if not separator or not colon or meeting != F08_MEETING_ID or not ref or not literal.strip() or ref in parsed:
            raise ValueError('Pinned fixture has a malformed or duplicate passage line.')
        parsed[ref] = literal
    if not parsed:
        raise ValueError('Pinned fixture has no passages.')
    return parsed


def validate_minutes_notes(data):
    """Validate the isolated F08 branch before opening any supplied source path."""
    required = {'meeting_id', 'meeting_date', 'cutoff', 'scope', 'sources', 'records', 'minutes_notes'}
    if not isinstance(data, dict) or set(data) != required:
        raise ValueError('Use the exact F08 synthetic minutes-notes fields.')
    if data['meeting_id'] != F08_MEETING_ID or data['meeting_date'] != '2026-09-14' or data['cutoff'] != '2026-09-14':
        raise ValueError('F08 synthetic meeting identity and dates are fixed.')
    if not required_text(data['scope']) or data['records'] != []:
        raise ValueError('F08 requires a review scope and empty carry-forward records.')
    notes = data['minutes_notes']
    if not isinstance(notes, dict) or set(notes) != {
        'schema_version', 'fixture_only', 'agenda_source_id', 'agenda_revision',
        'notes_source_id', 'notes_revision', 'entries'
    }:
        raise ValueError('Use the exact F08 minutes_notes fields.')
    if (type(notes['schema_version']) is not int or notes['schema_version'] != 1
            or notes['fixture_only'] is not True
            or notes['agenda_source_id'] != 'F08-AGENDA-01'
            or notes['notes_source_id'] != 'F08-NOTES-01'
            or notes['agenda_revision'] != F08_FIXTURES['F08-AGENDA-01']['revision']
            or notes['notes_revision'] != F08_FIXTURES['F08-NOTES-01']['revision']):
        raise ValueError('F08 source identity or revision does not match the pinned fixture manifest.')
    if not isinstance(data['sources'], list) or len(data['sources']) != 2:
        raise ValueError('F08 requires exactly the two pinned source fixtures.')
    sources = {}
    for source in data['sources']:
        if not isinstance(source, dict) or set(source) != {'id', 'path', 'sha256', 'link', 'revision', 'snapshot', 'access_scope'}:
            raise ValueError('F08 source must match the pinned fixture fields.')
        fixture = F08_FIXTURES.get(source['id'])
        if fixture is None or source['id'] in sources:
            raise ValueError('Unknown or duplicate F08 fixture ID.')
        # Compare the exact string before any filesystem operation. This rejects
        # arbitrary paths without even resolving or statting the supplied path.
        expected = {
            'path': str(fixture['path']), 'sha256': fixture['sha256'], 'link': fixture['link'],
            'revision': fixture['revision'], 'snapshot': fixture['snapshot'],
            'access_scope': 'approved_for_local_review',
        }
        if any(source.get(key) != value for key, value in expected.items()):
            raise ValueError('F08 source does not match its pinned fixture manifest.')
        sources[source['id']] = source
    if set(sources) != set(F08_FIXTURES):
        raise ValueError('Both distinct F08 fixtures are required.')
    agenda = _f08_lines(sources['F08-AGENDA-01'])
    note_lines = _f08_lines(sources['F08-NOTES-01'])
    entries = notes['entries']
    if not isinstance(entries, list):
        raise ValueError('F08 note entries must be a list.')
    used_notes = set()
    for entry in entries:
        if not isinstance(entry, dict) or set(entry) != {'agenda_item_ref', 'note_ref', 'kind'}:
            raise ValueError('Each F08 note entry must contain only its agenda ref, note ref and kind.')
        note_ref, agenda_ref, kind = entry['note_ref'], entry['agenda_item_ref'], entry['kind']
        if note_ref not in note_lines or note_ref in used_notes or kind not in F08_NOTE_KINDS:
            raise ValueError('F08 note reference is unknown, duplicated or has an unsupported kind.')
        if kind in {'attendance', 'off_agenda_concern'}:
            if agenda_ref is not None:
                raise ValueError('Attendance and off-agenda concerns must not be linked to an agenda item.')
        elif agenda_ref not in agenda:
            raise ValueError('F08 note entry must link to an existing agenda passage.')
        used_notes.add(note_ref)
    if used_notes != set(note_lines):
        raise ValueError('Every pinned notes passage must be mapped exactly once.')
    return agenda, note_lines, sources


def _f08_output_path(output):
    """Confine synthetic verbatim output to the canonical OS temp directory."""
    root = Path(tempfile.gettempdir()).resolve()
    requested = Path(output)
    resolved = requested.resolve()
    try:
        resolved.relative_to(root)
    except ValueError as error:
        raise ValueError('F08 output must be inside the canonical OS temporary directory.') from error
    if requested.suffix.lower() != '.docx' or resolved == root:
        raise ValueError('F08 output must be a DOCX file inside the canonical OS temporary directory.')
    return resolved


def build_minutes_notes(data, output):
    agenda, note_lines, sources = validate_minutes_notes(data)
    output = _f08_output_path(output)
    if output.exists():
        raise FileExistsError('Existing synthetic draft retained; do not overwrite it.')
    doc = Document()
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(11)
    style.paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25.4)
    def line(text):
        doc.add_paragraph(text, style='Normal')
    banner = doc.add_paragraph(style='Normal')
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    banner.paragraph_format.space_before = Pt(6)
    banner.paragraph_format.space_after = Pt(12)
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'FCE4D6')
    banner._p.get_or_add_pPr().append(shade)
    run = banner.add_run('SYNTHETIC TEST OUTPUT — NOT ACTUAL MINUTES — NOT FOR CIRCULATION')
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(156, 0, 6)
    line('Clinical Directors — fictional minutes draft | unapproved local test only')
    line('Meeting: F08-DEMO-01 | Fictional meeting date: 2026-09-14')
    line('Source-note mapping is structural only; it does not verify completeness, truth, attendance or authority.')
    agenda_source = sources['F08-AGENDA-01']
    notes_source = sources['F08-NOTES-01']
    line(f"Agenda source: {agenda_source['link']} | {agenda_source['revision']} | SHA-256: {agenda_source['sha256']}")
    line(f"Notes source: {notes_source['link']} | {notes_source['revision']} | SHA-256: {notes_source['sha256']}")
    by_agenda = {}
    off_agenda = []
    unlinked = []
    for entry in data['minutes_notes']['entries']:
        if entry['kind'] == 'off_agenda_concern':
            off_agenda.append(entry)
        elif entry['agenda_item_ref'] is None:
            unlinked.append(entry)
        else:
            by_agenda.setdefault(entry['agenda_item_ref'], []).append(entry)
    labels = {
        'attendance': 'Attendance source note — verify actual attendance and quorum.',
        'discussion': 'Discussion source note.',
        'reported_decision': 'Reported decision source note — authority unverified; not an authorised decision.',
        'action_candidate': 'Action candidate source note — not commissioned, assigned or issued.',
        'condition': 'Condition source note — verify linkage and wording.',
        'dissent': 'Dissent source note — preserve and verify attribution.',
        'held': 'Held source note — no decision or action inferred.',
        'unresolved': 'Unresolved source note — human confirmation required.',
    }
    line('Agenda items and linked source notes')
    for entry in unlinked:
        line(f"{labels[entry['kind']]} {notes_source['id']}#{F08_MEETING_ID}:{entry['note_ref']} — {note_lines[entry['note_ref']]}")
    for agenda_ref, agenda_text in agenda.items():
        line(f'Agenda {agenda_ref}: {agenda_text}')
        mapped = by_agenda.get(agenda_ref, [])
        if not mapped:
            line('UNRESOLVED — No source-linked note supplied; outcome remains to be confirmed.')
        for entry in mapped:
            note = notes_source['id']
            line(f"{labels[entry['kind']]} {note}#{F08_MEETING_ID}:{entry['note_ref']} — {note_lines[entry['note_ref']]}")
    line('Off-agenda concerns — unresolved')
    for entry in off_agenda:
        line(f"UNRESOLVED concern source note {notes_source['id']}#{F08_MEETING_ID}:{entry['note_ref']} — {note_lines[entry['note_ref']]}")
    line('The existing concern route remains available without a paper or proposed solution. Urgent concerns are not delayed for minutes preparation.')
    line('No register write, assignment, due date, publication, notification or closure is inferred.')
    doc.core_properties.author = doc.core_properties.last_modified_by = ''
    raw = BytesIO()
    doc.save(raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as handle:
        handle.write(raw.getvalue())
    return output


def select(data):
    if isinstance(data, dict) and 'minutes_notes' in data:
        validate_minutes_notes(data)
        return [], [], {source['id']: source for source in data['sources']}
    if not isinstance(data, dict) or set(data) != {'meeting_id', 'meeting_date', 'cutoff', 'scope', 'sources', 'records'}:
        raise ValueError('Use the documented meeting, scope, source and record fields.')
    if not all(required_text(data[k]) for k in ('meeting_id', 'scope')):
        raise ValueError('Meeting ID and review scope are required.')
    meeting, cutoff = iso_date(data['meeting_date']), iso_date(data['cutoff'])
    if cutoff > meeting:
        raise ValueError('Cut-off cannot follow the target meeting.')
    if not isinstance(data['sources'], list) or not data['sources'] or not isinstance(data['records'], list):
        raise ValueError('Source snapshots and records must be lists; a source is required.')
    sources = {}
    for source in data['sources']:
        if not isinstance(source, dict) or set(source) != {'id', 'path', 'sha256', 'link', 'revision', 'snapshot', 'access_scope'}:
            raise ValueError('Each source needs ID, local path, hash, link, revision, snapshot time and access scope.')
        if not all(required_text(v) for v in source.values()):
            raise ValueError('Source details must be explicit.')
        if source['id'] in sources:
            raise ValueError('Duplicate source ID.')
        if source['access_scope'] == 'restricted_pending':
            raise ValueError('Restricted source is not read or used; resolve access in the approved route first.')
        if source['access_scope'] != 'approved_for_local_review':
            raise ValueError('Use an explicit known access-scope declaration.')
        timestamp = datetime.fromisoformat(source['snapshot'])
        if timestamp.tzinfo is None or timestamp.astimezone(MEETING_TIMEZONE).date() > cutoff:
            raise ValueError('Source snapshot needs a timezone and must not follow cut-off.')
        path = Path(source['path'])
        if not path.is_file() or sha256(path.read_bytes()).hexdigest() != source['sha256']:
            raise ValueError('Source unavailable or changed; reconcile before drafting.')
        sources[source['id']] = source
    selected, excluded, ids = [], [], set()
    for row in data['records']:
        if not isinstance(row, dict) or set(row) != ROW_FIELDS:
            raise ValueError('Unexpected or missing record fields.')
        if any(not required_text(row[k]) for k in ROW_FIELDS - {'review_date', 'due_date', 'verified', 'closure_date', 'closure_authority', 'closure_evidence_ref'}):
            raise ValueError('Record text must be explicit; use Not recorded where appropriate.')
        identity = row['id'].strip().casefold()
        if identity in ids:
            raise ValueError('Duplicate record ID; reconcile sources before drafting.')
        ids.add(identity)
        if row['source_id'] not in sources or row['state'] not in STATES or row['type'] not in {'action', 'decision'}:
            raise ValueError('Unknown source, record type or selection state.')
        if type(row['verified']) is not bool:
            raise ValueError('Verification must be explicit true or false.')
        allowed = {'action': {'open', 'closure_pending', 'closed', 'candidate'},
                   'decision': {'follow_up', 'closure_pending', 'closed', 'candidate', 'no_follow_up'}}
        if row['state'] not in allowed[row['type']]:
            raise ValueError('Selection state conflicts with record type.')
        if row['review_status'] not in {'draft', 'source_reviewed', 'approved_record'}:
            raise ValueError('Unknown source-review status.')
        if row['type'] == 'action':
            if row['commissioning'] not in {'unclear', 'not_commissioned', 'commissioned'} or row['issue_readiness'] not in {'hold', 'ready', 'issued'}:
                raise ValueError('Action commissioning and issue readiness must be separate known states.')
            if row['issue_readiness'] in {'ready', 'issued'} and (row['commissioning'] != 'commissioned' or row['review_status'] == 'draft' or not evidenced_text(row['owner']) or not evidenced_text(row['authority']) or not (row['due_date'] or row['review_date'])):
                raise ValueError('Ready or issued action requires source review, a commission, evidenced authority, owner and agreed due or review date.')
        elif row['commissioning'] != 'not_applicable' or row['issue_readiness'] != 'not_applicable':
            raise ValueError('Decision cannot be assigned action-issue states.')
        if row['type'] == 'decision' and row['review_status'] == 'approved_record' and not evidenced_text(row['authority']):
            raise ValueError('Approved decision requires evidenced deciding authority.')
        review = iso_date(row['review_date']) if row['review_date'] is not None else None
        due = iso_date(row['due_date']) if row['due_date'] is not None else None
        closure = iso_date(row['closure_date']) if row['closure_date'] is not None else None
        if row['closure_acceptance'] not in {'unknown', 'reported_complete', 'accepted'}:
            raise ValueError('Closure acceptance must have a known state.')
        if claims_closed(row['status']) and row['state'] != 'closed':
            raise ValueError('Recorded closed status conflicts with unresolved selection state; reconcile the source.')
        if row['state'] == 'closed' and (
            row['closure_acceptance'] != 'accepted' or not closure or closure > cutoff
            or not evidenced_text(row['closure_authority'])
            or not known_source_locator(row['closure_evidence_ref'], sources, row['id'])
        ):
            raise ValueError('Accepted closure needs a separate accepted state, dated authority and source reference by the cut-off; reported completion is not closure.')
        if row['state'] != 'closed' and row['closure_acceptance'] == 'accepted':
            raise ValueError('Accepted closure conflicts with a non-closed selection state.')
        if row['closure_acceptance'] == 'reported_complete' and row['state'] != 'closure_pending':
            raise ValueError('Reported completion requires closure-pending review, not open or accepted closure.')
        if not row['verified'] or row['state'] == 'candidate':
            reason = 'Unverified or uncommissioned candidate; excluded from carry-forward'
        elif row['type'] == 'action' and row['review_status'] == 'draft':
            reason = 'Action source remains draft; excluded from action carry-forward'
        elif row['type'] == 'action' and row['commissioning'] != 'commissioned':
            reason = 'Commissioning not evidenced; excluded from action carry-forward'
        elif row['type'] == 'action' and row['issue_readiness'] != 'issued':
            # A descriptive owner field does not establish an accepted,
            # accountable owner. Keep every missing particular visible while
            # preserving the existing exclusion decision.
            particulars = []
            if not evidenced_text(row['owner']):
                particulars.append('accountable owner not evidenced/accepted in supplied snapshot')
            if not (due or review):
                particulars.append('agreed due/review date not recorded in supplied snapshot')
            reason = 'Action not issued; retain for preparation review only'
            if particulars:
                reason += ' — ' + '; '.join(particulars)
        elif row['type'] == 'decision' and row['review_status'] != 'approved_record':
            reason = 'Decision source not approved; excluded from decision carry-forward'
        elif row['state'] in {'closed', 'no_follow_up'}:
            reason = 'Recorded closed or no explicit follow-up'
        elif review and review > meeting and not (due and due <= meeting):
            reason = 'Review date is after the target meeting'
        else:
            selected.append(row)
            continue
        excluded.append((row, reason))
    return selected, excluded, sources


def build(data, output):
    if isinstance(data, dict) and 'minutes_notes' in data:
        return build_minutes_notes(data, output)
    selected, excluded, sources = select(data)
    output = Path(output)
    if output.suffix.lower() != '.docx':
        raise ValueError('Output must be DOCX.')
    if output.exists():
        raise FileExistsError('Existing draft retained. Reconcile changes in that draft; do not create routine numbered copies.')
    doc = Document()
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(11)
    style.paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25.4)
    def line(text):
        doc.add_paragraph(text, style='Normal')
    line('Clinical Directors — draft minutes')
    line('LOCAL REVIEW DRAFT — not issued. Attendance and new outcomes are not yet recorded.')
    line(f"Meeting: {data['meeting_id']} | Date: {data['meeting_date']} | Cut-off: {data['cutoff']}")
    line('Review scope: ' + data['scope'])
    line('Attendance, apologies and conflicts — to record')
    line('Name | descriptive service/delegate capacity | present/apology/absent/unconfirmed | relevant arrival/departure. Verify actual attendance; absence cannot be inferred from connection data. Record quorum and authority separately.')
    line('Confirmation of previous minutes — outcome not yet recorded')
    line('Prior outstanding business — source snapshot, not new outcomes')
    line('Carry-forward does not renew instructions, change deadlines or accept closure. Review included and omitted items against the sources. Source hashes establish integrity only; human verification must establish commissioning and authority.')
    line('The source-selection, access-scope and record-review values are supplied declarations. This generator does not authenticate permission, source authority or factual completeness.')
    for source in sources.values():
        line(f"Source {source['id']}: {source['link']} | Revision: {source['revision']} | Snapshot: {source['snapshot']} | SHA-256: {source['sha256']}")
    if not selected:
        line('No eligible rows in the supplied reviewed snapshot. This does not establish that no outstanding business exists.')
    for row in selected:
        line(f"{row['id']} — {row['type']}: {row['text']}")
        line(f"Recorded status: {row['status']} | Recorded owner (descriptive): {row['owner']} | Recorded due date: {row['due_date'] or 'Not recorded'} | Review eligibility: {row['review_date'] or 'Not recorded — review required'}")
        line('Recorded authority: ' + row['authority'])
        line(f"Source review: {row['review_status']} | Commissioning: {row['commissioning']} | Issue readiness: {row['issue_readiness']}")
        line('Conditions: ' + row['conditions'])
        line('Uncertainty/dissent: ' + row['dissent'])
        line('Dependencies: ' + row['dependencies'])
        line('Source: ' + row['source_id'])
        if row['state'] == 'closure_pending':
            line('Closure remains awaiting acceptance; reported completion is not accepted closure.')
    line('Preparation review — excluded source IDs to reconcile before circulation; exclusion is not a decision to drop work')
    for row, reason in excluded:
        detail = f" | Closure reference: {row['closure_evidence_ref']}" if row['state'] == 'closed' else ''
        line(f"{row['id']} | Source: {row['source_id']} | Recorded snapshot status: {row['status']} | {reason}{detail}")
        line('Recorded source text (review only; not carried forward): ' + row['text'])
        line(f"Recorded owner (descriptive): {row['owner']} | Original due date: {row['due_date'] or 'Not recorded'} | Review eligibility: {row['review_date'] or 'Not recorded — review required'}")
        line('Recorded authority: ' + row['authority'])
        line(f"Source review: {row['review_status']} | Commissioning: {row['commissioning']} | Issue readiness: {row['issue_readiness']}")
        line('Conditions: ' + row['conditions'])
        line('Uncertainty/dissent: ' + row['dissent'])
        line('Dependencies: ' + row['dependencies'])
    if not excluded:
        line('None excluded from the supplied snapshot. Source completeness still requires human review.')
    line('This meeting: items, decisions, actions, conditions and dissent — not yet recorded')
    line('Record actual outcomes and authority after the meeting. No automatic register writeback, assignments or notifications.')
    line('Next meeting — to confirm')
    doc.core_properties.author = doc.core_properties.last_modified_by = ''
    raw = BytesIO()
    doc.save(raw)
    with output.open('xb') as handle:
        handle.write(raw.getvalue())
    return output


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    try:
        build(json.loads(args.input.read_text()), args.output)
    except FileExistsError:
        parser.exit(1, 'Draft not prepared. Existing draft retained; reconcile changes in that file instead of creating a numbered copy.\n')
    except ValueError as error:
        parser.exit(1, f'Draft not prepared: {error} Originals retained.\n')
    except (TypeError, OSError) as error:
        parser.exit(1, f'Draft not prepared due to an input or file-access error ({type(error).__name__}). Check source access, integrity and input fields; originals retained.\n')
    print('Local review draft prepared; human source-to-output check required.')


if __name__ == '__main__':
    main()
