#!/usr/bin/env python3
"""Build the editable Word editions of the Clinical Directors governance suite.

The Markdown and CSV files remain the version-controlled source material. This
script creates consistently formatted DOCX editions for operational use and
adds plain-text Word content controls to the three automation templates.
"""
from __future__ import annotations

import csv
import hashlib
import os
import re
import shutil
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]

DOCS_WORD = ROOT / "docs" / "word"
TEMPLATES_WORD = ROOT / "templates" / "word"
CONTROL_MAP_PATH = ROOT / "templates" / "word-content-control-map.csv"

OUTPUTS = {
    "terms": DOCS_WORD / "Clinical_Directors_Terms_of_Reference_v0.1.docx",
    "protocol": DOCS_WORD / "Clinical_Directors_Operating_Protocol_v0.1.docx",
    "onboarding": DOCS_WORD / "Clinical_Directors_Member_Onboarding_Guide_v0.1.docx",
    "evaluation": DOCS_WORD / "Clinical_Directors_Evaluation_and_Improvement_Framework_v0.1.docx",
    "codex": DOCS_WORD / "Clinical_Directors_Existing_Workspace_Setup_Plan_v0.5.docx",
    "pack": TEMPLATES_WORD / "Clinical_Directors_Meeting_Pack_and_Agenda_Template.docx",
    "paper": TEMPLATES_WORD / "Clinical_Directors_Meeting_Paper_Template.docx",
    "minutes": TEMPLATES_WORD / "Clinical_Directors_Minutes_Template.docx",
}

SOURCES = {
    "terms": ROOT / "docs" / "34-clinical-directors-terms-of-reference-v0.1.md",
    "protocol": ROOT / "docs" / "35-clinical-directors-operating-protocol-v0.1.md",
    "onboarding": ROOT / "docs" / "47-clinical-directors-onboarding-guide-v0.1.md",
    "evaluation": ROOT / "docs" / "41-evaluation-framework-and-data-dictionary.md",
    "codex": ROOT / "docs" / "50-existing-workspace-setup-plan-v0.5.md",
}

# Palette aligned to the existing repository brief, without asserting an
# official Queensland Health corporate template.
NAVY = "17365D"
BLUE = "2F75B5"
MID_BLUE = "5B9BD5"
PALE_BLUE = "D9EAF7"
LIGHT_BLUE = "EAF2F8"
PALE_GREY = "F2F2F2"
MID_GREY = "D9E1F2"
DARK_GREY = "595959"
WHITE = "FFFFFF"
AMBER = "FFF2CC"
PALE_RED = "FCE4D6"

FONT = "Times New Roman"
FIXED_TIME = datetime(2026, 9, 1, 0, 0, 0, tzinfo=timezone.utc)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W15_NS = "http://schemas.microsoft.com/office/word/2012/wordml"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
PLACEHOLDER_RE = re.compile(r"\{\{([A-Za-z0-9_-]+)\}\}")

# Friendly visible prompts keep the templates readable while the underlying
# Word content-control tags remain stable for Power Automate and other tools.
CONTROL_DISPLAY = {
    "PaperTitle": "[Paper title]",
    "PaperID": "[Paper ID]",
    "MeetingID": "[Meeting ID]",
    "MeetingDate": "[Meeting date]",
    "MeetingDateTime": "[Date and time]",
    "LocationOrTeamsLink": "[Location or Teams link]",
    "Chair": "[Chair]",
    "Secretariat": "[Secretariat]",
    "PackVersion": "[Pack version]",
    "PackPublishedAt": "[Published date and time]",
    "SensitivityLabel": "[Classification]",
    "AgendaSeq": "[Seq]",
    "AgendaTime": "[Time]",
    "AgendaItemID": "[Item ID]",
    "AgendaPurpose": "[Purpose]",
    "AgendaDecisionMode": "[Decision mode]",
    "AgendaItemTitle": "[Item title]",
    "AgendaPresenter": "[Presenter]",
    "AgendaPaperReference": "[Paper]",
    "AgendaDispositionSought": "[Disposition sought]",
    "PreviousMinutesStatusAndLinkOrContent": "[Previous minutes status, link or content]",
    "ActionID": "[Action ID]",
    "ActionText": "[Action]",
    "ActionOwner": "[Owner]",
    "ActionDueReview": "[Due or review date]",
    "ActionStatus": "[Status]",
    "ActionException": "[Blocker or exception]",
    "ManifestPaperID": "[Paper ID]",
    "ManifestAgendaItemID": "[Item ID]",
    "ManifestPaperTitle": "[Title]",
    "ManifestSponsor": "[Sponsor]",
    "ManifestVersionStatus": "[Version or status]",
    "ManifestClassification": "[Classification]",
    "ManifestDocumentLink": "[Document or link]",
    "RestrictedReferencesOnly": "[Restricted references only]",
    "Author": "[Author]",
    "Sponsor": "[Sponsor]",
    "Purpose": "[Purpose]",
    "DecisionMode": "[Decision mode]",
    "VersionStatus": "[Version or status]",
    "SensitivityRoute": "[Classification or sensitivity route]",
    "DecisionSought": "[Decision, advice or assurance sought]",
    "Recommendation": "[Recommendation]",
    "ExecutiveSummary": "[Executive summary]",
    "Context": "[Context and current position]",
    "PersonAndEquityImpact": "[Patients, people, communities and equity impacts]",
    "EvidenceAndUncertainty": "[Evidence, data, assumptions and uncertainty]",
    "Options": "[Options and trade-offs]",
    "SafetyQuality": "[Clinical safety and quality implications]",
    "Workforce": "[Workforce and capability implications]",
    "Operational": "[Operational and interface implications]",
    "Financial": "[Financial and resource implications]",
    "LegalDigital": "[Legal, regulatory, privacy, records and digital implications]",
    "CulturalSafetyEquity": "[Cultural safety and equity implications]",
    "Communication": "[Communication and stakeholder implications]",
    "RisksAndControls": "[Material risks, controls, residual risk and owner]",
    "Consultation": "[Consultation and differing views]",
    "Implementation": "[Implementation, ownership and monitoring]",
    "NextSteps": "[Next steps following the decision]",
    "Appendices": "[Appendices or links]",
    "MinutesStatus": "[Minutes status]",
    "MinutesVersion": "[Minutes version]",
    "RecordingTranscriptionStatus": "[Recording and transcription status]",
    "MinutesClassification": "[Classification]",
    "AttendanceAndConflicts": "[Attendance, apologies, quorum and conflicts]",
    "PreviousMinutesConfirmation": "[Previous minutes confirmation]",
    "MaterialConsiderations": "[Material considerations]",
    "Disposition": "[Disposition or decision]",
    "Authority": "[Authority or decision-maker]",
    "UncertaintyAndDissent": "[Material uncertainty or dissent]",
    "FollowUp": "[Residual risk, monitoring, communication and return condition]",
    "NextMeeting": "[Next meeting date, time and expected business]",
    "CandidateID": "[Candidate ID]",
    "CandidateType": "[Type]",
    "CandidateDraftText": "[Draft minutes text]",
    "TranscriptTime": "[Transcript time]",
    "TranscriptSpeaker": "[Speaker]",
    "SourceExcerpt": "[Source excerpt]",
    "ReviewerDisposition": "[Reviewer status or correction]",
}


SOURCE_FIELD_OVERRIDES = {
    # Meeting pack
    ("MeetingPack", "MeetingID"): "MeetingSchedule.MeetingID",
    ("MeetingPack", "MeetingDateTime"): "MeetingSchedule.MeetingDateTime",
    ("MeetingPack", "LocationOrTeamsLink"): "MeetingSchedule.TeamsMeetingLink",
    ("MeetingPack", "PackVersion"): "DocumentMetadata.Version",
    ("MeetingPack", "PackPublishedAt"): "MeetingSchedule.PackPublishedAt",
    ("MeetingPack", "SensitivityLabel"): "MeetingSchedule.SensitivityRoute",
    ("MeetingPack", "AgendaSeq"): "Agenda.Sequence",
    ("MeetingPack", "AgendaTime"): "Agenda.TimeAllocationMinutes",
    ("MeetingPack", "AgendaItemID"): "Agenda.AgendaItemID",
    ("MeetingPack", "AgendaPurpose"): "Agenda.Purpose",
    ("MeetingPack", "AgendaDecisionMode"): "Agenda.DecisionMode",
    ("MeetingPack", "AgendaItemTitle"): "Agenda.Title",
    ("MeetingPack", "AgendaPresenter"): "Agenda.Presenter",
    ("MeetingPack", "AgendaPaperReference"): "Agenda.PaperLink",
    ("MeetingPack", "AgendaDispositionSought"): "Agenda.DecisionSought",
    ("MeetingPack", "PreviousMinutesStatusAndLinkOrContent"): "PriorMeeting.MinutesLink",
    ("MeetingPack", "ActionID"): "Action.ActionID",
    ("MeetingPack", "ActionText"): "Action.ActionText",
    ("MeetingPack", "ActionOwner"): "Action.Owner",
    ("MeetingPack", "ActionDueReview"): "Action.DueDate|ReviewDate",
    ("MeetingPack", "ActionStatus"): "Action.Status",
    ("MeetingPack", "ActionException"): "Action.BlockerOrException",
    ("MeetingPack", "ManifestPaperID"): "Paper.PaperID",
    ("MeetingPack", "ManifestAgendaItemID"): "Paper.AgendaItemID",
    ("MeetingPack", "ManifestPaperTitle"): "Paper.Title",
    ("MeetingPack", "ManifestSponsor"): "Paper.Sponsor",
    ("MeetingPack", "ManifestVersionStatus"): "Paper.Version|Status",
    ("MeetingPack", "ManifestClassification"): "Paper.ConfidentialityRoute",
    ("MeetingPack", "ManifestDocumentLink"): "Paper.AcceptedVersionLink|PaperDocumentLink",
    ("MeetingPack", "RestrictedReferencesOnly"): "Agenda.RestrictedReference",
    # Meeting paper
    ("Paper", "PaperTitle"): "Paper.Title",
    ("Paper", "MeetingID"): "Paper.RequestedMeetingID",
    ("Paper", "MeetingDate"): "MeetingSchedule.MeetingDateTime",
    ("Paper", "VersionStatus"): "Paper.Version|Status",
    ("Paper", "SensitivityRoute"): "Paper.ConfidentialityRoute",
    ("Paper", "Context"): "Paper.ContextCurrentPosition",
    ("Paper", "EvidenceAndUncertainty"): "Paper.EvidenceDataUncertainty",
    ("Paper", "Options"): "Paper.OptionsTradeOffs",
    ("Paper", "SafetyQuality"): "Paper.MaterialImplications.ClinicalSafetyQuality",
    ("Paper", "Workforce"): "Paper.MaterialImplications.WorkforceCapability",
    ("Paper", "Operational"): "Paper.MaterialImplications.OperationalInterfaces",
    ("Paper", "Financial"): "Paper.MaterialImplications.FinancialResources",
    ("Paper", "LegalDigital"): "Paper.MaterialImplications.LegalPrivacyRecordsDigital",
    ("Paper", "CulturalSafetyEquity"): "Paper.MaterialImplications.CulturalSafetyEquity",
    ("Paper", "Communication"): "Paper.MaterialImplications.CommunicationStakeholders",
    ("Paper", "Consultation"): "Paper.ConsultationAndDifferingViews",
    ("Paper", "Implementation"): "Paper.ImplementationAndMonitoring",
    ("Paper", "Appendices"): "Paper.AttachmentsManifest",
    # Minutes and reviewer provenance
    ("Minutes", "MeetingID"): "MeetingSchedule.MeetingID",
    ("Minutes", "MeetingDateTime"): "MeetingSchedule.MeetingDateTime",
    ("Minutes", "MinutesStatus"): "DocumentMetadata.DocumentStatus",
    ("Minutes", "MinutesVersion"): "DocumentMetadata.Version",
    ("Minutes", "RecordingTranscriptionStatus"): "MeetingSchedule.RecordingTranscriptionPlanned|TranscriptSourceLink|RecordingLink",
    ("Minutes", "MinutesClassification"): "DocumentMetadata.SensitivityRoute",
    ("Minutes", "AgendaItemID"): "Agenda.AgendaItemID",
    ("Minutes", "AgendaItemTitle"): "Agenda.Title",
    ("Minutes", "Purpose"): "Agenda.Purpose",
    ("Minutes", "DecisionMode"): "Agenda.DecisionMode",
    ("Minutes", "MaterialConsiderations"): "Decision.Rationale",
    ("Minutes", "Disposition"): "Decision.DecisionText|Agenda.Disposition",
    ("Minutes", "Authority"): "Decision.DecisionMakerRole|AuthorityReference",
    ("Minutes", "UncertaintyAndDissent"): "Decision.EvidenceAndUncertainty|SignificantDissent",
    ("Minutes", "ActionID"): "Action.ActionID",
    ("Minutes", "ActionText"): "Action.ActionText",
    ("Minutes", "ActionOwner"): "Action.Owner",
    ("Minutes", "ActionDueReview"): "Action.DueDate|ReviewDate",
    ("Minutes", "ActionStatus"): "Action.Status",
    ("Minutes", "FollowUp"): "Decision.ResidualRiskAndMonitoring",
    ("Minutes", "CandidateDraftText"): "AIVerification.DraftText",
    ("Minutes", "TranscriptTime"): "AIVerification.TimestampDisplay",
    ("Minutes", "TranscriptSpeaker"): "AIVerification.SpeakerLabel",
    ("Minutes", "ReviewerDisposition"): "AIVerification.ReviewStatus|CorrectedText",
}


def humanise_control_tag(tag: str) -> str:
    """Return a concise visible label without changing the automation tag."""
    if tag in CONTROL_DISPLAY:
        return CONTROL_DISPLAY[tag]
    words = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", tag)
    words = words.replace(" ID", " ID").replace(" AI", " AI")
    return f"[{words[0].upper()}{words[1:] if len(words) > 1 else ''}]"


def set_run_font(run, name: str = FONT, size: float | None = None, bold: bool | None = None,
                 italic: bool | None = None, colour: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if colour:
        run.font.color.rgb = RGBColor.from_string(colour)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = tr_pr.find(qn("w:cantSplit"))
    if cant is None:
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)


def set_paragraph_border(paragraph, bottom: str | None = None, top: str | None = None,
                         colour: str = MID_BLUE, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge, value in (("top", top), ("bottom", bottom)):
        if value:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), value)
            el.set(qn("w:sz"), str(size))
            el.set(qn("w:space"), "1")
            el.set(qn("w:color"), colour)
            p_bdr.append(el)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def set_keep_lines(paragraph, value: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepLines"))
    if keep is None:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)
    keep.set(qn("w:val"), "1" if value else "0")


def add_field(paragraph, instruction: str, fallback: str = "1") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(XML_SPACE, "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, colour=DARK_GREY)


def add_header_footer(section, short_title: str, status: str, version: str) -> None:
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"CHHHS  |  {short_title.upper()}")
    set_run_font(r, size=8.5, bold=True, colour=BLUE)
    set_paragraph_border(p, bottom="single", colour=MID_BLUE, size=6)

    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{status}  |  {version}  |  Page ")
    set_run_font(r, size=8.5, colour=DARK_GREY)
    add_field(p, "PAGE", "1")
    r = p.add_run(" of ")
    set_run_font(r, size=8.5, colour=DARK_GREY)
    add_field(p, "NUMPAGES", "1")


def configure_page(section, landscape: bool = False, narrow: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
    margin = 1.45 if narrow else 1.75
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(margin)
    section.right_margin = Cm(margin)


def configure_styles(doc: Document, compact: bool = False) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10 if compact else 10.5)
    normal.paragraph_format.space_after = Pt(4.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    for name, size, colour, before, after in [
        ("Title", 24, NAVY, 0, 8),
        ("Subtitle", 12, BLUE, 0, 10),
        ("Heading 1", 15, NAVY, 12, 4),
        ("Heading 2", 12.5, BLUE, 9, 3),
        ("Heading 3", 11, NAVY, 7, 2),
        ("Heading 4", 10.5, BLUE, 5, 2),
    ]:
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(colour)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(10 if compact else 10.5)
        style.paragraph_format.space_after = Pt(2.5)

    if "Template Prompt" not in styles:
        style = styles.add_style("Template Prompt", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Template Prompt"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(9)
    style.font.italic = True
    style.font.color.rgb = RGBColor.from_string(DARK_GREY)
    style.paragraph_format.space_after = Pt(2)

    if "Table Text" not in styles:
        style = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Table Text"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(8.5 if compact else 9)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def apply_dense_protocol_layout(doc: Document) -> None:
    """Compact a long operational guide without compromising readability."""
    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    for style_name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = doc.styles[style_name]
        style.font.size = Pt(9.5)
        style.paragraph_format.space_after = Pt(1.5)
    for style_name in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = doc.styles[style_name]
        style.paragraph_format.space_before = Pt(max(4, (style.paragraph_format.space_before or Pt(0)).pt - 2))
        style.paragraph_format.space_after = Pt(2)
    for section in doc.sections:
        section.top_margin = Cm(1.35)
        section.bottom_margin = Cm(1.35)



def set_core_properties(doc: Document, title: str, subject: str) -> None:
    props = doc.core_properties
    props.title = title
    props.subject = subject
    props.author = "Fictional service"
    props.last_modified_by = "CHHHS Clinical Hub Automation"
    props.category = "Clinical governance"
    props.comments = "Generated from the version-controlled CHHHS Clinical Hub repository."
    props.keywords = "CHHHS; Clinical Directors; governance; meeting"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1


def new_document(title: str, short_title: str, status: str, version: str,
                 landscape: bool = False, compact: bool = False, narrow: bool = False) -> Document:
    doc = Document()
    configure_styles(doc, compact=compact)
    section = doc.sections[0]
    configure_page(section, landscape=landscape, narrow=narrow)
    add_header_footer(section, short_title=short_title, status=status, version=version)
    set_core_properties(doc, title=title, subject=short_title)
    return doc


def add_title_block(doc: Document, title: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title)
    set_run_font(r, size=24, bold=True, colour=NAVY)
    set_paragraph_border(p, bottom="single", colour=MID_BLUE, size=10)
    if subtitle:
        p = doc.add_paragraph(style="Subtitle")
        r = p.add_run(subtitle)
        set_run_font(r, size=12, colour=BLUE)


def add_inline_markdown(paragraph, text: str, base_size: float | None = None,
                        base_colour: str | None = None) -> None:
    # Lightweight inline parser for links, emphasis, inline code and placeholders.
    token_re = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\{\{[A-Za-z0-9_-]+\}\}|\*\*.+?\*\*|`.+?`|(?<!\*)\*[^*]+?\*(?!\*))")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            r = paragraph.add_run(text[position:match.start()])
            set_run_font(r, size=base_size, colour=base_colour)
        token = match.group(0)
        if token.startswith("["):
            label, url = token[1:-1].split("](", 1)
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            r = paragraph.add_run(label)
            set_run_font(r, size=base_size, colour=BLUE)
            r.underline = True
            link.append(r._r)
            paragraph._p.append(link)
        elif token.startswith("{{"):
            r = paragraph.add_run(token)
            set_run_font(r, size=base_size, italic=True, colour=BLUE)
        elif token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            set_run_font(r, size=base_size, bold=True, colour=base_colour)
        elif token.startswith("`"):
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, name="Aptos Mono", size=(base_size or 9.5), colour=DARK_GREY)
        else:
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, size=base_size, italic=True, colour=base_colour)
        position = match.end()
    if position < len(text):
        r = paragraph.add_run(text[position:])
        set_run_font(r, size=base_size, colour=base_colour)


def add_callout(doc: Document, text: str, fill: str = LIGHT_BLUE, border: str = MID_BLUE,
                bold_prefix: bool = False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=150, end=150)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "5")
        el.set(qn("w:color"), border)
        borders.append(el)
    p = cell.paragraphs[0]
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_after = Pt(0)
    add_inline_markdown(p, text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def clear_cell(cell) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "Table Text"
    p.paragraph_format.space_after = Pt(0)


def write_cell(cell, text: str, bold: bool = False, colour: str | None = None,
               align: WD_ALIGN_PARAGRAPH | None = None, size: float | None = None) -> None:
    clear_cell(cell)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    add_inline_markdown(p, text, base_size=size, base_colour=colour)
    if bold:
        for run in p.runs:
            run.bold = True
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def style_header_row(row, font_size: float = 8.5) -> None:
    set_row_repeat(row)
    set_row_cant_split(row)
    for cell in row.cells:
        set_cell_shading(cell, NAVY)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                set_run_font(run, size=font_size, bold=True, colour=WHITE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=90, bottom=90, start=80, end=80)


def add_table(doc: Document, rows: Sequence[Sequence[str]], widths_cm: Sequence[float] | None = None,
              first_col_labels: bool = False, compact: bool = False) -> object:
    if not rows:
        return doc.add_table(rows=0, cols=0)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        set_row_cant_split(row)
        for j, text in enumerate(row_data):
            write_cell(row.cells[j], text, size=8 if compact else 9)
            if widths_cm:
                set_cell_width(row.cells[j], widths_cm[j])
        if i == 0:
            style_header_row(row, font_size=8 if compact else 8.5)
        elif first_col_labels:
            set_cell_shading(row.cells[0], LIGHT_BLUE)
            for run in row.cells[0].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_metadata_table(doc: Document, pairs: Sequence[tuple[str, str]], columns: int = 2,
                       widths_cm: Sequence[float] | None = None) -> object:
    if columns not in (1, 2):
        raise ValueError("columns must be 1 or 2")
    cols = columns * 2
    rows_count = (len(pairs) + columns - 1) // columns
    table = doc.add_table(rows=rows_count, cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    index = 0
    for row in table.rows:
        set_row_cant_split(row)
        for group in range(columns):
            label_cell = row.cells[group * 2]
            value_cell = row.cells[group * 2 + 1]
            if index < len(pairs):
                label, value = pairs[index]
                write_cell(label_cell, label, bold=True, colour=NAVY, size=9)
                set_cell_shading(label_cell, LIGHT_BLUE)
                write_cell(value_cell, value, size=9)
            else:
                write_cell(label_cell, "", size=9)
                write_cell(value_cell, "", size=9)
            if widths_cm:
                set_cell_width(label_cell, widths_cm[group * 2])
                set_cell_width(value_cell, widths_cm[group * 2 + 1])
            index += 1
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def split_markdown_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_markdown_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_markdown_table(doc: Document, block: list[str], landscape: bool = False) -> None:
    rows = [split_markdown_row(block[0])] + [split_markdown_row(x) for x in block[2:]]
    ncols = len(rows[0])
    # Normalise malformed rows conservatively.
    rows = [(r + [""] * ncols)[:ncols] for r in rows]
    if ncols == 2:
        widths = [4.1, 12.9] if not landscape else [5.0, 20.5]
        add_table(doc, rows, widths_cm=widths, first_col_labels=True)
    elif ncols == 5:
        widths = [3.5, 9.2, 3.6, 2.7, 6.4] if landscape else [2.7, 6.2, 2.5, 2.1, 3.6]
        add_table(doc, rows, widths_cm=widths, compact=True)
    else:
        total = 25.4 if landscape else 17.0
        add_table(doc, rows, widths_cm=[total / ncols] * ncols, compact=ncols >= 5)


def add_numbered_paragraph(doc: Document, number: str, text: str, level: int = 0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65 + 0.55 * level)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(2.5)
    r = p.add_run(f"{number}. ")
    set_run_font(r, bold=True)
    add_inline_markdown(p, text)


def add_bullet_paragraph(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet 2" if level else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2.5)
    add_inline_markdown(p, text)


def add_body_paragraph(doc: Document, text: str, indent: bool = False) -> None:
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.85)
        p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(4.5)
    # Preserve Markdown hard breaks.
    parts = re.split(r"\s{2,}\n", text)
    for idx, part in enumerate(parts):
        add_inline_markdown(p, part.replace("\n", " ").strip())
        if idx < len(parts) - 1:
            p.add_run().add_break()


def convert_markdown_document(source: Path, output: Path, short_title: str,
                              status: str, version: str, subtitle: str,
                              landscape: bool = False, compact: bool = False, narrow: bool = False,
                              dense: bool = False, body_font_size: float | None = None,
                              page_break_before: tuple[str, ...] = ()) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    doc = new_document(title, short_title, status, version, landscape=landscape,
                       compact=compact, narrow=narrow)
    if dense:
        apply_dense_protocol_layout(doc)
    if body_font_size is not None:
        doc.styles["Normal"].font.size = Pt(body_font_size)
    add_title_block(doc, title, subtitle)

    i = 1
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith(">"):
            block: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            add_callout(doc, " ".join(block))
            continue

        heading_match = re.match(r"^(#{2,6})\s+(.+)$", stripped)
        if heading_match:
            md_level = len(heading_match.group(1))
            level = min(max(md_level - 1, 1), 4)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_markdown(p, heading_match.group(2))
            set_keep_with_next(p)
            if heading_match.group(2) in page_break_before:
                p.paragraph_format.page_break_before = True
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            block = [lines[i], lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_markdown_table(doc, block, landscape=landscape)
            continue

        bullet = re.match(r"^(\s*)[-*]\s+(.+)$", line)
        if bullet:
            level = 1 if len(bullet.group(1).replace("\t", "    ")) >= 2 else 0
            add_bullet_paragraph(doc, bullet.group(2).rstrip(), level=level)
            i += 1
            continue

        numbered = re.match(r"^(\s*)(\d+)\.\s+(.+)$", line)
        if numbered:
            level = 1 if len(numbered.group(1).replace("\t", "    ")) >= 2 else 0
            add_numbered_paragraph(doc, numbered.group(2), numbered.group(3).rstrip(), level=level)
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            set_paragraph_border(p, bottom="single", colour=MID_GREY, size=6)
            i += 1
            continue

        # Collect a normal paragraph until a structural line or blank.
        paragraph_lines = [line.rstrip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            nst = nxt.strip()
            if not nst:
                break
            if (
                nst.startswith("#")
                or nst.startswith(">")
                or re.match(r"^\s*[-*]\s+", nxt)
                or re.match(r"^\s*\d+\.\s+", nxt)
                or (nst.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]))
                or nst == "---"
            ):
                break
            paragraph_lines.append(nxt.rstrip())
            i += 1
        indent = bool(paragraph_lines[0].startswith("   "))
        add_body_paragraph(doc, "\n".join(x.strip() for x in paragraph_lines), indent=indent)

    save_final(doc, output, wrap_controls=False)


def add_template_prompt(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Template Prompt")
    p.add_run(text)


def add_entry_box(doc: Document, placeholder: str, min_lines: int = 2, fill: str = "FFFFFF") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=120, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_markdown(p, placeholder)
    for _ in range(max(0, min_lines - 1)):
        p.add_run().add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


F06_TEMPLATE_FIELDS = {
    'MeetingID': ('SourceSnapshot.meeting_id', 'Meeting occurrence identity'),
    'MeetingDateTime': ('SelectedOccurrence.date|timezone', 'Date-only stays date-only; do not invent a time or booking'),
    'DraftRevision': ('Generation.output_revision', 'Separate output revision from native file version'),
    'OccurrenceScope': ('SourceSnapshot.scope', 'Fictional scope and unissued status; no authority promotion'),
    'PaperLink': ('SelectedPaper.web_url|label', 'Rich-text hyperlink; generate explicit relationship, not a URL-only text claim'),
    'PaperIdentity': ('SelectedPaper.unique_id|native_version', 'Keep selected source version and identity'),
    'SourceHistory': ('SourceSnapshot.prior_meeting_id|antecedent_source', 'Preserve prior minutes/register references; local paths are not cloud links'),
    **{f'AgendaSection{i}': (f'AgendaContract.sections[{i-1}].text', 'Reviewed seven-section wording; do not prefill new outcomes') for i in range(1, 8)},
    **{f'ReviewHold{i}': (f'SourceSnapshot.records[{i-1}]', 'Full ID, unknown status, source_text, history, original owner/due/review and response reference; no filtering, assignment or closure') for i in range(1, 4)},
    'SourceLink': ('SourceSnapshot.native_item_url|label', 'Rich-text hyperlink; generate explicit relationship'),
    'SourceProvenance': ('SourceSnapshot.source_id|source_revision|native_version|modified|included_ids', 'Preserve included records and source cutoff'),
    'TemplateProvenance': ('Template.contract_id|version|sha256', 'Native template version remains unresolved until uploaded/read back'),
    'RunProvenance': ('Generation.flow_id|run_id|created_at|output_revision', 'No fabricated cloud run; local specimens must say local preparation only'),
    'OptionalContent': ('Generation.optional_content', 'May be empty; no extra record or invented outcome'),
}

F06_TEMPLATE_PROMPTS = {
    'MeetingID': 'Meeting occurrence reference',
    'MeetingDateTime': 'Meeting date and timezone; time only if confirmed',
    'DraftRevision': 'Unissued draft revision',
    'OccurrenceScope': 'Preparation scope and unresolved booking or participation details',
    'PaperLink': 'Link to the selected paper',
    'PaperIdentity': 'Selected paper identity and version',
    'SourceHistory': 'Prior minutes/register references and corrections; preserve original history',
    'AgendaSection1': 'Welcome, attendance, apologies and interests only as observed',
    'AgendaSection2': 'Early paper or guest contribution, if selected; no attendance presumed',
    'AgendaSection3': 'Executive update, if supplied',
    'AgendaSection4': 'Previous minutes, corrections and unresolved accuracy questions',
    'AgendaSection5': 'Source-linked review questions and unresolved concerns',
    **{f'ReviewHold{i}': 'Original ID, status, full wording and history; original owner, dates and response' for i in range(1, 4)},
    'AgendaSection6': 'Existing actions and returned responses; retain original IDs and dates',
    'AgendaSection7': 'Trial/noting items and closing recap; new outcomes and next booking remain blank',
    'SourceLink': 'Link to the source record',
    'SourceProvenance': 'Source identity, revision, cut-off and included records',
    'TemplateProvenance': 'Template version and reference',
    'RunProvenance': 'Generation reference; generation is not approval',
    'OptionalContent': 'Additional context, only if needed; otherwise leave blank',
}


def build_meeting_pack(output: Path | None = None) -> None:
    """One-time maintained F06 template authoring; no operational population/API."""
    source = ROOT / 'templates' / 'clinical-directors-meeting-pack.md'
    lines = source.read_text(encoding='utf-8').splitlines()
    declared = PLACEHOLDER_RE.findall('\n'.join(lines))
    if len(declared) != len(set(declared)) or set(declared) != set(F06_TEMPLATE_FIELDS):
        raise ValueError('F06 maintained source/control contract mismatch')
    doc = new_document(lines[0].removeprefix('# '), 'Agenda draft', 'Unissued template',
                       'F06 Word v1', landscape=False, compact=True, narrow=False)
    doc.styles['Normal'].font.size = Pt(11)
    doc.core_properties.subject = 'F06 seven-section agenda template; local preparation, not cloud output'
    doc.core_properties.comments = 'Contract F06-word-agenda-seven-section-v1. No property bindings or native repeating-section claim.'
    # Keep this selected template free of legacy callout-table/border markup.
    # Word property order matters to SDK validation even when a renderer opens it.
    for paragraph in doc.sections[0].header.paragraphs:
        for border in list(paragraph._p.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
    doc.add_heading(lines[0].removeprefix('# '), level=0)
    for line in lines[1:]:
        line = line.strip()
        if not line:
            continue
        match = PLACEHOLDER_RE.search(line)
        if match:
            tag = match.group(1)
            prefix = line[:match.start()].replace('**', '').strip()
            if prefix:
                p = doc.add_paragraph(prefix)
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0)
            block = OxmlElement('w:sdt')
            props = OxmlElement('w:sdtPr')
            for name in ('alias', 'tag'):
                node = OxmlElement(f'w:{name}'); node.set(qn('w:val'), tag); props.append(node)
            identity = OxmlElement('w:id'); identity.set(qn('w:val'), str(600 + declared.index(tag))); props.append(identity)
            if tag not in {'PaperLink', 'SourceLink'}:
                text = OxmlElement('w:text'); text.set(qn('w:multiLine'), '1'); props.append(text)
            content = OxmlElement('w:sdtContent')
            paragraph = doc.add_paragraph()
            paragraph.add_run(f'[{F06_TEMPLATE_PROMPTS[tag]}]')
            content.append(paragraph._p)
            block.append(props); block.append(content)
            doc._element.body.insert(-1, block)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('> '):
            doc.add_paragraph(line[2:]).runs[0].italic = True
        else:
            add_inline_markdown(doc.add_paragraph(), line)
    save_final(doc, output or OUTPUTS['pack'], wrap_controls=False)


def build_paper() -> None:
    title = "{{PaperTitle}}"
    doc = new_document("Clinical Directors Meeting Paper Template", "Meeting paper", "Template", "v0.1")
    add_title_block(doc, title, "Decision-oriented paper for the Clinical Directors Meeting")
    add_metadata_table(
        doc,
        [
            ("Paper ID", "{{PaperID}}"),
            ("Meeting / date", "{{MeetingID}} / {{MeetingDate}}"),
            ("Author", "{{Author}}"),
            ("Sponsor", "{{Sponsor}}"),
            ("Purpose", "{{Purpose}}"),
            ("Decision mode", "{{DecisionMode}}"),
            ("Version / status", "{{VersionStatus}}"),
            ("Classification", "{{SensitivityRoute}}"),
        ],
        columns=2,
        widths_cm=[3.1, 5.4, 3.1, 5.4],
    )
    add_callout(doc, "State the decision, advice, assurance, problem-solving or escalation sought on page 1. Keep the paper proportionate to significance, risk and uncertainty.")

    sections = [
        ("1. Decision, advice or assurance sought", "State exactly what the forum or decision-maker is being asked to do.", "{{DecisionSought}}", 3),
        ("2. Recommendation", "State the preferred course and why.", "{{Recommendation}}", 3),
        ("3. Executive summary", "Summarise the issue, recommendation, material consequences and urgency.", "{{ExecutiveSummary}}", 4),
        ("4. Context and current position", "Describe the relevant history, current state and why the matter requires this forum.", "{{Context}}", 4),
        ("5. Patients, people, communities and equity", "Address patient, family, carer, workforce, First Nations, cultural safety, rural and remote, access and equity impacts where relevant.", "{{PersonAndEquityImpact}}", 4),
        ("6. Evidence, data and uncertainty", "Distinguish evidence, data, professional judgement, assumptions, limitations and unresolved uncertainty.", "{{EvidenceAndUncertainty}}", 4),
        ("7. Options and trade-offs", "Set out realistic alternatives, including the status quo, and compare benefits, harms, feasibility, resources and reversibility.", "{{Options}}", 5),
    ]
    for heading, prompt, placeholder, lines in sections:
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(heading)
        add_template_prompt(doc, prompt)
        add_entry_box(doc, placeholder, min_lines=lines)

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("8. Material implications")
    implications = [
        ("Clinical safety and quality", "{{SafetyQuality}}"),
        ("Workforce and capability", "{{Workforce}}"),
        ("Operational and interfaces", "{{Operational}}"),
        ("Financial and resources", "{{Financial}}"),
        ("Legal, regulatory, privacy, records and digital", "{{LegalDigital}}"),
        ("Cultural safety and equity", "{{CulturalSafetyEquity}}"),
        ("Communication and stakeholders", "{{Communication}}"),
    ]
    for heading, placeholder in implications:
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(heading)
        add_entry_box(doc, placeholder, min_lines=2)

    remaining = [
        ("9. Risks and controls", "Identify the material risks, controls, residual risk and accountable owner.", "{{RisksAndControls}}", 4),
        ("10. Consultation and differing views", "Record who was consulted, material differences in view and any perspectives still missing.", "{{Consultation}}", 3),
        ("11. Implementation, ownership and monitoring", "Set out accountable ownership, sequence, resources, dependencies, measures and review arrangements.", "{{Implementation}}", 4),
        ("12. Next steps following the decision", "State the immediate actions, communication, escalation or return-to-forum requirements.", "{{NextSteps}}", 3),
        ("Appendices", "Attach only material needed to support informed consideration or assurance.", "{{Appendices}}", 1),
    ]
    for heading, prompt, placeholder, lines in remaining:
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(heading)
        add_template_prompt(doc, prompt)
        add_entry_box(doc, placeholder, min_lines=lines)

    save_final(doc, OUTPUTS["paper"], wrap_controls=True)


def build_minutes() -> None:
    title = "Clinical Directors Meeting Minutes"
    doc = new_document(title, "Minutes", "Template", "v0.1")
    add_title_block(doc, title, "Decision, assurance, escalation and accountability record")
    add_metadata_table(
        doc,
        [
            ("Meeting ID", "{{MeetingID}}"),
            ("Date and time", "{{MeetingDateTime}}"),
            ("Chair", "{{Chair}}"),
            ("Secretariat", "{{Secretariat}}"),
            ("Status", "{{MinutesStatus}}"),
            ("Version", "{{MinutesVersion}}"),
            ("Recording / transcription", "{{RecordingTranscriptionStatus}}"),
            ("Classification", "{{MinutesClassification}}"),
        ],
        columns=2,
        widths_cm=[3.1, 5.4, 3.1, 5.4],
    )
    add_callout(doc, "These minutes are a concise governance record, not a transcript. Draft text remains subject to accountable human review and the approved confirmation process.")

    for heading, placeholder, lines in [
        ("Attendance, apologies, quorum and conflicts", "{{AttendanceAndConflicts}}", 3),
        ("Confirmation of previous minutes", "{{PreviousMinutesConfirmation}}", 2),
    ]:
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(heading)
        add_entry_box(doc, placeholder, min_lines=lines)

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Items")
    item_table = doc.add_table(rows=1, cols=1)
    item_table.style = "Table Grid"
    item_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = item_table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_markdown(p, "{{AgendaItemID}} — {{AgendaItemTitle}}", base_colour=NAVY)
    for run in p.runs:
        run.bold = True
    add_metadata_table(doc, [("Purpose", "{{Purpose}}"), ("Decision mode", "{{DecisionMode}}")], columns=2,
                       widths_cm=[2.8, 5.7, 2.8, 5.7])

    item_sections = [
        ("Material considerations", "Summarise only what is necessary to understand the disposition and governance trail.", "{{MaterialConsiderations}}", 4),
        ("Disposition / decision / advice / assurance / escalation", "Record the explicit outcome. Do not use 'discussed' as an approval status.", "{{Disposition}}", 3),
        ("Authority / decision-maker", "Identify the actual decision-maker and authority or delegation, where applicable.", "{{Authority}}", 2),
        ("Material uncertainty or dissent", "Record significant unresolved uncertainty or dissent where relevant, without unnecessary attribution.", "{{UncertaintyAndDissent}}", 3),
    ]
    for heading, prompt, placeholder, lines in item_sections:
        p = doc.add_paragraph(style="Heading 2")
        p.add_run(heading)
        add_template_prompt(doc, prompt)
        add_entry_box(doc, placeholder, min_lines=lines)

    p = doc.add_paragraph(style="Heading 2")
    p.add_run("Actions")
    action_headers = ["Action ID", "Action", "Owner", "Due / review", "Status"]
    action_row = ["{{ActionID}}", "{{ActionText}}", "{{ActionOwner}}", "{{ActionDueReview}}", "{{ActionStatus}}"]
    add_table(doc, [action_headers, action_row], widths_cm=[2.4, 7.0, 3.0, 2.8, 2.2], compact=True)

    p = doc.add_paragraph(style="Heading 2")
    p.add_run("Residual risk, monitoring, communication and return condition")
    add_entry_box(doc, "{{FollowUp}}", min_lines=3)

    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Next meeting")
    add_entry_box(doc, "{{NextMeeting}}", min_lines=2, fill=PALE_GREY)

    # Reviewer appendix uses a landscape section because the provenance table is wide.
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(section, landscape=True, narrow=True)
    add_header_footer(section, short_title="Minutes — reviewer provenance appendix", status="Reviewer-only template", version="v0.1")
    p = doc.add_paragraph(style="Heading 1")
    p.add_run("Reviewer-only AI provenance appendix")
    add_callout(
        doc,
        "Remove this appendix from the approved circulation copy unless local records advice requires retention in the same document. Preserve the controlled verification record.",
        fill=AMBER,
        border="BF9000",
    )
    provenance_headers = ["Candidate ID", "Type", "Draft minutes text", "Transcript time", "Speaker", "Source excerpt", "Reviewer status / correction"]
    provenance_row = [
        "{{CandidateID}}", "{{CandidateType}}", "{{CandidateDraftText}}", "{{TranscriptTime}}",
        "{{TranscriptSpeaker}}", "{{SourceExcerpt}}", "{{ReviewerDisposition}}",
    ]
    add_table(doc, [provenance_headers, provenance_row], widths_cm=[2.5, 2.1, 5.0, 2.8, 2.8, 6.5, 4.8], compact=True)

    save_final(doc, OUTPUTS["minutes"], wrap_controls=True)


def qn_lxml(local: str) -> str:
    return f"{{{W_NS}}}{local}"


def make_sdt(tag: str, placeholder_text: str, rpr: etree._Element | None) -> etree._Element:
    sdt = etree.Element(qn_lxml("sdt"))
    sdt_pr = etree.SubElement(sdt, qn_lxml("sdtPr"))
    alias = etree.SubElement(sdt_pr, qn_lxml("alias"))
    alias.set(qn_lxml("val"), tag)
    tag_el = etree.SubElement(sdt_pr, qn_lxml("tag"))
    tag_el.set(qn_lxml("val"), tag)
    ident = etree.SubElement(sdt_pr, qn_lxml("id"))
    # Stable positive 31-bit identifier.
    stable_id = int(hashlib.sha256(tag.encode("utf-8")).hexdigest()[:8], 16) & 0x7FFFFFFF
    ident.set(qn_lxml("val"), str(stable_id))
    etree.SubElement(sdt_pr, qn_lxml("text"))
    etree.SubElement(sdt_pr, qn_lxml("showingPlcHdr"))
    colour = etree.SubElement(sdt_pr, f"{{{W15_NS}}}color")
    colour.set(f"{{{W15_NS}}}val", BLUE)

    content = etree.SubElement(sdt, qn_lxml("sdtContent"))
    run = etree.SubElement(content, qn_lxml("r"))
    if rpr is not None:
        run.append(deepcopy(rpr))
    text = etree.SubElement(run, qn_lxml("t"))
    text.text = humanise_control_tag(tag)
    return sdt


def wrap_placeholders_in_part(path: Path) -> int:
    parser = etree.XMLParser(remove_blank_text=False)
    tree = etree.parse(str(path), parser)
    root = tree.getroot()
    changed = 0
    runs = list(root.xpath(".//w:r", namespaces={"w": W_NS}))
    for run in runs:
        texts = run.xpath("./w:t", namespaces={"w": W_NS})
        if not texts:
            continue
        full = "".join(t.text or "" for t in texts)
        if not PLACEHOLDER_RE.search(full):
            continue
        parent = run.getparent()
        if parent is None:
            continue
        rpr = run.find("w:rPr", namespaces={"w": W_NS})
        idx = parent.index(run)
        pos = 0
        replacements: list[etree._Element] = []
        for match in PLACEHOLDER_RE.finditer(full):
            if match.start() > pos:
                new_run = etree.Element(qn_lxml("r"))
                if rpr is not None:
                    new_run.append(deepcopy(rpr))
                new_text = etree.SubElement(new_run, qn_lxml("t"))
                new_text.text = full[pos:match.start()]
                if new_text.text.startswith(" ") or new_text.text.endswith(" "):
                    new_text.set(XML_SPACE, "preserve")
                replacements.append(new_run)
            tag = match.group(1)
            replacements.append(make_sdt(tag, match.group(0), rpr))
            pos = match.end()
        if pos < len(full):
            new_run = etree.Element(qn_lxml("r"))
            if rpr is not None:
                new_run.append(deepcopy(rpr))
            new_text = etree.SubElement(new_run, qn_lxml("t"))
            new_text.text = full[pos:]
            if new_text.text.startswith(" ") or new_text.text.endswith(" "):
                new_text.set(XML_SPACE, "preserve")
            replacements.append(new_run)
        for node in replacements:
            parent.insert(idx, node)
            idx += 1
        parent.remove(run)
        changed += 1
    tree.write(str(path), xml_declaration=True, encoding="UTF-8", standalone="yes")
    return changed


def deterministic_repack(source: Path, destination: Path, wrap_controls: bool) -> None:
    with tempfile.TemporaryDirectory(prefix="chhhs_docx_") as td:
        work = Path(td) / "unzipped"
        work.mkdir(parents=True)
        with zipfile.ZipFile(source, "r") as zf:
            zf.extractall(work)

        if wrap_controls:
            for part in [work / "word" / "document.xml", *sorted((work / "word").glob("header*.xml")), *sorted((work / "word").glob("footer*.xml"))]:
                if part.exists():
                    wrap_placeholders_in_part(part)

        destination.parent.mkdir(parents=True, exist_ok=True)
        tmp_out = destination.with_suffix(".tmp.docx")
        if tmp_out.exists():
            tmp_out.unlink()
        with zipfile.ZipFile(tmp_out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            for file_path in sorted(p for p in work.rglob("*") if p.is_file()):
                arcname = file_path.relative_to(work).as_posix()
                info = zipfile.ZipInfo(arcname, date_time=(1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                zf.writestr(info, file_path.read_bytes())
        tmp_out.replace(destination)


def save_final(doc: Document, output: Path, wrap_controls: bool) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="chhhs_docx_save_") as td:
        intermediate = Path(td) / "intermediate.docx"
        doc.save(intermediate)
        deterministic_repack(intermediate, output, wrap_controls=wrap_controls)


def extract_content_control_tags(path: Path) -> list[str]:
    with zipfile.ZipFile(path, "r") as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    tags: list[str] = []
    for node in root.xpath(".//w:sdtPr/w:tag", namespaces={"w": W_NS}):
        value = node.get(qn_lxml("val"))
        if value and value not in tags:
            tags.append(value)
    return tags


def control_map_attributes(template: str, tag: str) -> tuple[str, str, str, str]:
    """Return logical repeat group, source field, required flag and note."""
    if template == "MeetingPack" and tag in F06_TEMPLATE_FIELDS:
        source, note = F06_TEMPLATE_FIELDS[tag]
        return "", source, "No" if tag == "OptionalContent" else "Yes", note
    repeat_group = ""
    source = f"{template}.{tag}"
    required = "No"
    note = ""

    if template == "MeetingPack":
        if tag.startswith("Agenda"):
            repeat_group = "AgendaRows"
            source = f"Agenda.{tag.removeprefix('Agenda')}"
        elif tag.startswith("Action"):
            repeat_group = "OpenActionRows"
            source = f"Action.{tag.removeprefix('Action')}"
        elif tag.startswith("Manifest"):
            repeat_group = "PaperManifestRows"
            source = f"Paper.{tag.removeprefix('Manifest')}"
        else:
            source = f"Meeting.{tag}"
        required = "Yes" if tag in {
            "MeetingID", "MeetingDateTime", "Chair", "AgendaItemID", "AgendaPurpose",
            "AgendaItemTitle", "AgendaDispositionSought",
        } else "No"
    elif template == "Paper":
        source = f"Paper.{tag}"
        required = "Yes" if tag in {
            "PaperTitle", "PaperID", "MeetingID", "Purpose", "DecisionMode",
            "DecisionSought", "Recommendation",
        } else "No"
        if tag in {
            "DecisionSought", "Recommendation", "ExecutiveSummary", "Context",
            "PersonAndEquityImpact", "EvidenceAndUncertainty", "Options", "SafetyQuality",
            "Workforce", "Operational", "Financial", "LegalDigital", "CulturalSafetyEquity",
            "Communication", "RisksAndControls", "Consultation", "Implementation",
            "NextSteps", "Appendices",
        }:
            note = "Allow carriage returns"
    elif template == "Minutes":
        provenance = {
            "CandidateID", "CandidateType", "CandidateDraftText", "TranscriptTime",
            "TranscriptSpeaker", "SourceExcerpt", "ReviewerDisposition",
        }
        item_fields = {
            "AgendaItemID", "AgendaItemTitle", "Purpose", "DecisionMode",
            "MaterialConsiderations", "Disposition", "Authority", "UncertaintyAndDissent",
            "FollowUp",
        }
        if tag.startswith("Action"):
            repeat_group = "ActionRows"
            source = f"Action.{tag.removeprefix('Action')}"
        elif tag in provenance:
            repeat_group = "ProvenanceRows"
            source = f"AIVerification.{tag}"
        elif tag in item_fields:
            repeat_group = "ItemRows"
            source = f"MinutesItem.{tag}"
        else:
            source = f"Meeting.{tag}"
        required = "Yes" if tag in {
            "MeetingID", "MeetingDateTime", "MinutesStatus", "AgendaItemID",
            "AgendaItemTitle", "Disposition",
        } else "No"

    source = SOURCE_FIELD_OVERRIDES.get((template, tag), source)
    if repeat_group and not note:
        note = "Logical template row; repeating-section behaviour requires tenant validation"
    return repeat_group, source, required, note


def write_control_map() -> None:
    template_files = [
        ("MeetingPack", OUTPUTS["pack"]),
        ("Paper", OUTPUTS["paper"]),
        ("Minutes", OUTPUTS["minutes"]),
    ]
    CONTROL_MAP_PATH.parent.mkdir(parents=True, exist_ok=True)
    with CONTROL_MAP_PATH.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow([
            "Template", "ControlName", "ControlType", "RepeatGroup",
            "SourceField", "Required", "Notes",
        ])
        for template, path in template_files:
            for tag in extract_content_control_tags(path):
                repeat_group, source, required, note = control_map_attributes(template, tag)
                kind = "RichText" if template == "MeetingPack" and tag in {"PaperLink", "SourceLink"} else "PlainText"
                writer.writerow([template, tag, kind, repeat_group, source, required, note])


def write_readmes() -> None:
    DOCS_WORD.mkdir(parents=True, exist_ok=True)
    TEMPLATES_WORD.mkdir(parents=True, exist_ok=True)
    (DOCS_WORD / "README.md").write_text(
        """# Word governance documents

These editable DOCX documents are generated from the version-controlled Markdown sources by `scripts/build_word_documents.py`. The Markdown and machine-readable files remain the source of truth until an approved CHHHS/Queensland Health document-control process supersedes them.

- `Clinical_Directors_Terms_of_Reference_v0.1.docx`
- `Clinical_Directors_Operating_Protocol_v0.1.docx`
- `Clinical_Directors_Member_Onboarding_Guide_v0.1.docx`
- `Clinical_Directors_Evaluation_and_Improvement_Framework_v0.1.docx`
- `Clinical_Directors_Existing_Workspace_Setup_Plan_v0.5.docx`
""",
        encoding="utf-8",
    )
    (TEMPLATES_WORD / "README.md").write_text(
        """# Word operational templates

These DOCX templates are editable and contain tagged plain-text Word content controls for deterministic or human-assisted population. Template rows are usable manually. Repeating-section behaviour must be validated in the CHHHS Microsoft 365 tenant before production use.

- `Clinical_Directors_Meeting_Pack_and_Agenda_Template.docx`
- `Clinical_Directors_Meeting_Paper_Template.docx`
- `Clinical_Directors_Minutes_Template.docx`

The exact control tags, logical row groups and proposed source fields are recorded in `../word-content-control-map.csv`.
""",
        encoding="utf-8",
    )


def main() -> int:
    DOCS_WORD.mkdir(parents=True, exist_ok=True)
    TEMPLATES_WORD.mkdir(parents=True, exist_ok=True)

    convert_markdown_document(
        SOURCES["terms"], OUTPUTS["terms"], "Terms of reference", "Working draft", "v0.1",
        "Interim terms of reference for local review and approval",
    )
    convert_markdown_document(
        SOURCES["protocol"], OUTPUTS["protocol"], "Operating protocol", "Working draft", "v0.1",
        "Practical rules for the meeting cycle, governance records and automation",
        compact=True,
        narrow=True,
        dense=True,
    )
    convert_markdown_document(
        SOURCES["onboarding"], OUTPUTS["onboarding"], "Member onboarding", "Working guide", "v0.1",
        "A concise guide for members, presenters and invited attendees",
    )
    convert_markdown_document(
        SOURCES["evaluation"], OUTPUTS["evaluation"], "Evaluation framework", "Working draft", "v0.1",
        "Structure, process, outcome and balancing measures for iterative improvement",
        landscape=True,
        compact=True,
        narrow=True,
    )
    convert_markdown_document(
        SOURCES["codex"], OUTPUTS["codex"], "Codex setup plan", "Working draft", "v0.5",
        "Clinical Leadership (CHHHS) | Clinical Directors Meeting",
        compact=False,
        narrow=False,
        dense=False,
        body_font_size=11,
        page_break_before=("3. Keep the governance that matters", "6. Leave three useful records"),
    )
    build_meeting_pack()
    build_paper()
    build_minutes()
    write_control_map()
    write_readmes()

    for path in OUTPUTS.values():
        if not path.exists() or path.stat().st_size < 10_000:
            raise RuntimeError(f"Word output missing or unexpectedly small: {path}")
        print(path.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
